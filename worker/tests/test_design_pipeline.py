"""Module 3 end to end on synthetic pictures with known ground truth (tests/design_samples.py).

Each sample is analysed once per session. The thresholds are what the current pipeline achieves, with
a little room; a regression in OCR, typography, shapes, tables or the reconstruction shows here.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from otk_worker.design import assets

from conftest_converter import needs_chromium, needs_lo

pytestmark = pytest.mark.skipif(not (assets.fonts_dir() / "fonts.json").is_file() or assets.model_path("realesr-general-x4v3.onnx") is None,
                                reason="bundled fonts or models missing: run scripts/fetch_fonts.py and scripts/fetch_models.py")

LANGS = {"poster": ["eng", "hin"], "certificate": ["eng"], "scan": ["eng"], "small": ["eng"]}


@pytest.fixture(scope="session")
def analysed(tmp_path_factory) -> dict[str, tuple[Path, dict, dict]]:
    import design_samples

    from otk_worker.converter import ocr
    from otk_worker.design import pipeline

    if "hin" not in ocr.languages_available():
        pytest.skip("Tesseract Hindi data not installed")
    root = tmp_path_factory.mktemp("design")
    built = design_samples.build(root / "in")
    out = {}
    for name, (path, truth) in built.items():
        proj = root / f"proj-{name}"
        scene = pipeline.analyze(path, proj, langs=LANGS[name])
        out[name] = (proj, scene, truth)
    return out


def _same_color(a: str, b: str, tol: int = 12) -> bool:
    return sum(abs(int(a[i:i + 2], 16) - int(b[i:i + 2], 16)) for i in (1, 3, 5)) <= tol


def _text_offsets(scene: dict, a, b, reach: int = 10) -> list[str]:
    """For a failure message: where each text layer of picture `b` sits relative to `a`.

    "name dx,dy left%" = the shift (px) that best lines the layer's ink up, and how much of the
    difference is left after shifting (0% = only moved; high = different glyphs, size or font).
    """
    import numpy as np

    A = 255 - np.asarray(a.convert("L"), dtype=np.int32)
    B = 255 - np.asarray(b.convert("L"), dtype=np.int32)
    H, W = A.shape
    out = []
    for lyr in scene["layers"]:
        if lyr["type"] != "text":
            continue
        x, y, w, h = (int(round(v)) for v in lyr["box"])
        x, y = max(x, reach), max(y, reach)
        w, h = min(w, W - reach - x), min(h, H - reach - y)
        if w <= 0 or h <= 0:
            continue
        ref = A[y:y + h, x:x + w]
        errs = {(dx, dy): int(np.abs(B[y + dy:y + dy + h, x + dx:x + dx + w] - ref).sum())
                for dy in range(-reach, reach + 1) for dx in range(-reach, reach + 1)}
        (dx, dy), best = min(errs.items(), key=lambda kv: kv[1])
        out.append(f"{lyr['name'][:24]!r} {dx:+d},{dy:+d} left {100 * best / max(errs[0, 0], 1):.0f}%")
    return out


def _regions(r: dict) -> list[tuple]:
    return [(g["x"], g["y"], g["w"], g["h"]) for g in r["regions"]]


def _pdf_text(pdf: Path, width: int) -> list[str]:
    """For a failure message: the fonts LibreOffice used and where each text span landed (page px)."""
    import fitz

    with fitz.open(pdf) as doc:
        page = doc[0]
        k = width / page.rect.width
        out = ["fonts: " + ", ".join(f"{f[3]} ({f[1]}, {f[2]})" for f in page.get_fonts())]
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for sp in line["spans"]:
                    if sp["text"].strip():
                        x0, y0, x1, y1 = (round(v * k) for v in sp["bbox"])
                        out.append(f"{sp['text'][:20]!r} {sp['font']} {sp['size'] * k:.1f}px #{sp['color']:06x} "
                                   f"[{x0},{y0},{x1 - x0},{y1 - y0}]")
    return out


def test_every_text_line_is_read_and_styled(analysed):
    total = {"n": 0, "family": 0, "weight": 0, "size": 0, "color": 0, "align": 0}
    misses = []
    for name, (_, scene, truth) in analysed.items():
        texts = [l for l in scene["layers"] if l["type"] == "text"]
        for t in (e for e in truth["elements"] if e["kind"] == "text"):
            total["n"] += 1
            want = t["text"].replace("\n", " ").strip()
            got = next((l for l in texts if l["text"].replace("\n", " ").strip() == want), None)
            if got is None:
                misses.append(f"{name}: {want!r} not read exactly")
                continue
            st = got["style"]
            total["family"] += st["family"] == t["family"]
            total["weight"] += st["weight"] == t["weight"] and st["italic"] == t["italic"]
            total["size"] += abs(st["size"] - t["size"]) / t["size"] <= 0.05
            total["color"] += _same_color(st["color"].lower(), t["color"].lower())
            total["align"] += st["align"] == t["align"]
    assert not misses, misses
    n = total["n"]
    assert n == 20
    # Measured: 19/20 for family (Inter vs Roboto on 14 px bold text), weight, size and colour; 20/20 alignment.
    for k in ("family", "weight", "size", "color"):
        assert total[k] >= n - 2, (k, total)
    assert total["align"] == n, total


def test_poster_shapes_graphics_photo_and_table(analysed):
    _, scene, _ = analysed["poster"]
    layers = scene["layers"]
    assert layers[0]["role"] == "background"
    shapes = [l for l in layers if l["type"] == "shape"]
    button = next(l for l in shapes if l["shape"] == "rounded")
    assert button["fill"] == "#1e5aa8" and button["radius"] == pytest.approx(24, abs=3)
    assert any(l["shape"] == "ellipse" and l["fill"] == "#ffd7c2" for l in shapes)
    outline = next(l for l in shapes if l["shape"] == "rect" and l["stroke"])
    assert outline["fill"] is None and outline["strokeWidth"] == pytest.approx(5, abs=0.5)
    rule = next(l for l in shapes if l["shape"] == "line")
    assert rule["stroke"] == "#8b1e3f" and rule["strokeWidth"] == pytest.approx(4, abs=0.5)
    vectors = [l for l in layers if l["type"] == "vector"]
    assert len(vectors) == 2 and all(v["paths"] for v in vectors)
    assert any(v["colors"] == ["#f2a900"] for v in vectors)                                # the star
    # The badge: a native green circle, with its white tick traced as a graphic on top.
    badge = next(l for l in shapes if l["shape"] == "ellipse" and l["fill"] == "#2e8b57")
    tick = next(v for v in vectors if v["colors"] == ["#ffffff"])
    assert layers.index(tick) > layers.index(badge)
    photos = [l for l in layers if l["type"] == "image" and l.get("kind") == "photo"]
    assert len(photos) == 1 and photos[0]["box"][:2] == pytest.approx([60, 270], abs=2)
    table = next(l for l in layers if l["type"] == "table")
    assert len(table["rowHeights"]) == 4 and len(table["colWidths"]) == 3
    cells = {(c["row"], c["col"]): c for c in table["cells"]}
    assert [cells[(0, c)]["text"] for c in range(3)] == ["Day", "Stage", "Headliner"]
    assert cells[(1, 2)]["text"] == "The Night Owls"
    assert all(cells[(0, c)]["weight"] == 700 for c in range(3))
    assert all(cells[(r, c)]["weight"] == 400 for r in range(1, 4) for c in range(3))
    assert cells[(0, 0)]["fill"] == "#e8e2d6" and cells[(1, 0)]["fill"] is None
    assert table["style"]["family"] == "Roboto"


def test_certificate_frames_and_signature_lines(analysed):
    _, scene, _ = analysed["certificate"]
    shapes = [l for l in scene["layers"] if l["type"] == "shape"]
    frames = sorted((l for l in shapes if l["shape"] == "rect" and l["stroke"]), key=lambda l: -l["box"][2])
    assert [round(f["strokeWidth"]) for f in frames[:2]] == [6, 2]
    assert frames[0]["box"][:2] == pytest.approx([24, 24], abs=1)
    lines = [l for l in shapes if l["shape"] == "line"]
    assert len(lines) == 2
    captions = {l["text"]: l["style"]["align"] for l in scene["layers"] if l["type"] == "text"}
    assert captions["Director"] == "center" and captions["Date"] == "center"


def test_scan_is_straightened_and_its_table_read(analysed):
    _, scene, _ = analysed["scan"]
    # The sample is turned 1.6° counter-clockwise; straightening turns it back (clockwise is negative).
    assert scene["source"]["rotation"] == pytest.approx(-1.6, abs=0.3)
    table = next(l for l in scene["layers"] if l["type"] == "table")
    cells = {(c["row"], c["col"]): c["text"] for c in table["cells"]}
    assert cells[(1, 1)] == "1,240" and cells[(2, 2)] == "41"


def test_small_screenshot_uses_super_resolution_and_finds_the_controls(analysed):
    _, scene, _ = analysed["small"]
    assert scene["source"]["upscale"] >= 2
    shapes = [l for l in scene["layers"] if l["type"] == "shape"]
    assert any(l["shape"] == "rect" and l["fill"] == "#243b6b" for l in shapes)            # header bar at the edge
    field = next(l for l in shapes if l["shape"] == "rounded" and l["stroke"])
    assert field["fill"] == "#ffffff" and field["stroke"] == "#c8cfdb"                       # 1 px outline, white inside
    assert any(l["shape"] == "rounded" and l["fill"] == "#2f6feb" for l in shapes)


def test_unreadable_script_is_left_as_picture(analysed, tmp_path):
    from otk_worker.design import pipeline

    proj, _, _ = analysed["poster"]
    src = next(proj.parent.glob("in/poster.png"))
    scene = pipeline.analyze(src, tmp_path / "eng-only", langs=["eng"])
    assert len(scene["unreadable"]) == 1
    x, y, w, h = scene["unreadable"][0]["box"]
    assert 300 < x < 400 and 800 < y < 830
    assert not any(l["type"] == "text" and "#" in l["text"] for l in scene["layers"])
    assert any("could not be read" in n for n in scene["notes"])


def test_cache_makes_a_second_analysis_identical(analysed):
    from otk_worker.design import pipeline

    proj, scene, _ = analysed["certificate"]
    src = next(proj.parent.glob("in/certificate.png"))
    again = pipeline.analyze(src, proj, langs=["eng"])
    a, b = dict(scene), dict(again)
    a.pop("stats"), b.pop("stats")
    assert a == b
    assert not any(p.suffix == ".pkl" for p in (proj / "cache").iterdir())       # plain PNG and JSON only


@needs_chromium
def test_rebuilt_design_looks_like_the_picture(analysed):
    from otk_worker.design import verify

    from PIL import Image

    results, detail = {}, {}
    for name, (proj, scene, _) in analysed.items():
        r = verify.accuracy(scene, proj)
        results[name] = (r["ssim"], len(r["regions"]))
        if r["ssim"] < 0.98 or len(r["regions"]) > 3:
            detail[name] = (_regions(r), _text_offsets(scene, Image.open(proj / "assets" / "prepared.png"),
                                                       Image.open(proj / "cache" / "render" / "drawn.png")))
    for name, (ssim, regions) in results.items():
        assert ssim >= 0.98, (results, detail)
        assert regions <= 3, (results, detail)


@needs_chromium
@needs_lo
def test_office_exports_look_like_the_design(analysed, tmp_path):
    from PIL import Image

    from otk_worker.design import project, verify

    import sys

    for name in ("poster", "certificate"):
        proj, scene, _ = analysed[name]
        if sys.platform == "win32":
            # LibreOffice on Windows sees installed fonts only: install them as a user would (Export → Install
            # the fonts). On Linux render_office points fontconfig at the design's font files instead.
            assert project.install_fonts(scene)["installed"]
        drawn = verify.render_scene(scene, proj)
        W, H = scene["page"]["width"], scene["page"]["height"]
        for fmt in ("pptx", "docx"):
            out = tmp_path / f"{name}.{fmt}"
            project.export(proj, scene, fmt, out)
            img = verify.render_office(out, W, H, scene=scene)
            r = verify.compare(drawn, img)
            if not (r["ssim"] >= 0.975 and len(r["regions"]) <= 2):
                pytest.fail("\n".join([f"{name} {fmt}: SSIM {r['ssim']}, regions {_regions(r)}", "text offsets:",
                                        *_text_offsets(scene, drawn, img), "LibreOffice's PDF:",
                                        *_pdf_text(out.parent / f"{out.stem}-render" / f"{out.stem}.pdf", W)]))
        # HTML in Chromium.
        from otk_worker.converter import chromium
        from otk_worker.design import layout

        html = project.export(proj, scene, "html", tmp_path / f"{name}.html")["path"]
        k, _ = layout.page_scale(scene, max_in=1000)
        chromium.print_html(Path(html), tmp_path / f"{name}-html.pdf", page_size_pt=(W * k, H * k), allow_dir=tmp_path)
        r = verify.compare(drawn, verify._pdf_page_image(tmp_path / f"{name}-html.pdf", W, H))
        assert r["ssim"] >= 0.975 and not r["regions"], (name, "html", r["ssim"], r["regions"])
        assert isinstance(Image.open(proj / "assets" / "background.png").size, tuple)


def test_exports_are_structurally_editable(analysed, tmp_path):
    import zipfile
    from xml.etree import ElementTree as ET

    from docx import Document
    from pptx import Presentation

    from otk_worker.design import fonts_out, project

    proj, scene, _ = analysed["poster"]
    paths = {fmt: Path(project.export(proj, scene, fmt, tmp_path / f"poster.{fmt}")["path"]) for fmt in project.FORMATS}

    # PowerPoint: one slide, a text box per text layer with the Office font names, a native table.
    prs = Presentation(str(paths["pptx"]))
    slide = prs.slides[0]
    texts = {sh.text_frame.text: sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text}
    assert "Summer Music Festival" in texts
    run = texts["Summer Music Festival"].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Montserrat ExtraBold"
    tables = [sh for sh in slide.shapes if sh.has_table]
    assert len(tables) == 1 and tables[0].table.cell(1, 2).text == "The Night Owls"
    groups = [sh for sh in slide.shapes if sh.shape_type == 6]     # GROUP: traced graphics
    assert len(groups) == 2

    # Word: text boxes, a floating table, embedded fonts that are real TrueType files once de-obfuscated.
    doc = Document(str(paths["docx"]))
    assert doc.tables[0].cell(0, 2).text == "Headliner"
    body = doc.element.xml
    assert "Summer Music Festival" in body and "wpg:wgp" in body
    with zipfile.ZipFile(paths["docx"]) as z:
        ft = ET.fromstring(z.read("word/fontTable.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
        emb = ft.findall(".//w:embedRegular", ns) + ft.findall(".//w:embedBold", ns)
        assert len(emb) >= 5
        rels = z.read("word/_rels/fontTable.xml.rels").decode()
        key = emb[0].get("{%s}fontKey" % ns["w"])
        rid = emb[0].get("{%s}id" % ns["r"])
        target = rels.split(f'Id="{rid}"')[1].split('Target="')[1].split('"')[0]
        raw = fonts_out.obfuscate(z.read(f"word/{target}"), key)
        assert raw[:4] in (b"\x00\x01\x00\x00", b"true")
        assert "w:embedTrueTypeFonts" in z.read("word/settings.xml").decode()

    # SVG: one group per layer, live text, embedded fonts.
    svg = ET.parse(paths["svg"]).getroot()
    groups = svg.findall("{http://www.w3.org/2000/svg}g")
    assert len(groups) == len(scene["layers"])
    assert "Summer Music Festival" in paths["svg"].read_text(encoding="utf-8")
    assert "@font-face{font-family:'Noto Sans Devanagari'" in paths["svg"].read_text(encoding="utf-8")

    # HTML: editable text and cells.
    html = paths["html"].read_text(encoding="utf-8")
    assert html.count('contenteditable="true"') >= 7 + 12

    # Project file: opens to the same scene.
    again = project.unpack(paths["otkd"], tmp_path / "reopened")
    assert json.dumps(again, sort_keys=True) == json.dumps(scene, sort_keys=True)


def _kinds(scene: dict) -> dict:
    shapes = [l for l in scene["layers"] if l["type"] == "shape"]
    return {
        "rounded": [l for l in shapes if l["shape"] == "rounded"],
        "outline": [l for l in shapes if l["shape"] == "rect" and l["stroke"] and not l["fill"]],
        "ellipse": [l for l in shapes if l["shape"] == "ellipse"],
        "line": [l for l in shapes if l["shape"] == "line"],
        "vector": [l for l in scene["layers"] if l["type"] == "vector"],
        "photo": [l for l in scene["layers"] if l["type"] == "image" and l.get("kind") == "photo"],
        "table": [l for l in scene["layers"] if l["type"] == "table"],
        "text": [l["text"].replace("\n", " ") for l in scene["layers"] if l["type"] == "text"],
    }


def test_jpeg_compressed_poster(analysed, tmp_path):
    """JPEG ringing and chroma bleeding must not turn shapes into noisy graphics or leave slivers."""
    from PIL import Image

    from otk_worker.design import pipeline

    proj, _, truth = analysed["poster"]
    src = tmp_path / "poster.jpg"
    Image.open(next(proj.parent.glob("in/poster.png"))).convert("RGB").save(src, quality=70)
    scene = pipeline.analyze(src, tmp_path / "proj", langs=["eng", "hin"])
    k = _kinds(scene)
    assert len(k["rounded"]) == 1 and len(k["outline"]) == 1 and len(k["ellipse"]) == 2 and len(k["line"]) == 1
    assert k["outline"][0]["strokeWidth"] == pytest.approx(5, abs=0.6)
    assert k["line"][0]["strokeWidth"] == pytest.approx(4, abs=0.6)
    assert len(k["vector"]) == 2 and len(k["photo"]) == 1 and len(k["table"]) == 1
    wanted = [e["text"].replace("\n", " ") for e in truth["elements"] if e["kind"] == "text"]
    assert sorted(k["text"]) == sorted(wanted)


def test_large_picture_is_analysed_at_working_size(analysed, tmp_path):
    """A 3000 x 4000 picture (a phone photo of a poster) is analysed at 2250 x 3000; photos keep full detail."""
    from PIL import Image

    from otk_worker.design import pipeline

    proj, _, truth = analysed["poster"]
    src = tmp_path / "big.png"
    Image.open(next(proj.parent.glob("in/poster.png"))).convert("RGB").resize((3000, 4000), Image.Resampling.LANCZOS).save(src)
    scene = pipeline.analyze(src, tmp_path / "proj", langs=["eng", "hin"])
    assert (scene["page"]["width"], scene["page"]["height"]) == (2250, 3000)
    assert any("analysed at 2250 × 3000" in n for n in scene["notes"])
    k = _kinds(scene)
    s = 2250 / 1200
    assert len(k["rounded"]) == 1 and len(k["outline"]) == 1 and len(k["ellipse"]) == 2 and len(k["line"]) == 1
    assert k["outline"][0]["strokeWidth"] == pytest.approx(5 * s, abs=1.2)
    assert k["line"][0]["strokeWidth"] == pytest.approx(4 * s, abs=1.2)
    assert len(k["vector"]) == 2 and len(k["table"]) == 1 and not scene["unreadable"]
    with Image.open(tmp_path / "proj" / k["photo"][0]["asset"]) as im:
        assert im.width >= 0.98 * 520 * 2.5      # cropped from the full-resolution pixels
    wanted = [e["text"].replace("\n", " ") for e in truth["elements"] if e["kind"] == "text"]
    assert sorted(k["text"]) == sorted(wanted)

"""Converter jobs: passwords, never overwriting, cancel and resume, merging, PDF/A proof and NOT-PDFA naming."""

from __future__ import annotations

import json
import threading
from pathlib import Path

import pikepdf
import pytest
from conftest_converter import copy_sample, needs_gs, needs_lo, needs_tess, needs_verapdf

from otk_worker.converter import pdfa, runner
from otk_worker.converter.steps import STEPS, load_all


def batch(files, target, out: Path, tmp_path: Path, notify=None, cancel=None, **extra) -> dict:
    params = {"jobId": extra.pop("jobId", "t1"), "files": files, "target": target, "outputDir": str(out),
              "jobsDir": str(tmp_path / "jobs"), "options": extra.pop("options", {}), **extra}
    return runner.run_batch(params, notify or (lambda p: None), cancel)


def test_passwords_wrong_missing_right(samples, tmp_path):
    out = tmp_path / "out"
    f = str(samples["pdf_encrypted"])
    res = batch([{"path": f}], "txt", out, tmp_path, jobId="p1")
    assert res["results"][0]["status"] == "needs_password"
    res = batch([{"path": f, "password": "wrong"}], "txt", out, tmp_path, jobId="p2")
    assert res["results"][0]["status"] == "needs_password"
    assert "not correct" in res["results"][0]["message"]
    res = batch([{"path": f, "password": "secret-123"}], "txt", out, tmp_path, jobId="p3")
    r = res["results"][0]
    assert r["status"] == "done", r
    assert "Quarterly Report" in Path(r["output"]).read_text(encoding="utf-8")
    report = json.loads(Path(r["reportJson"]).read_text(encoding="utf-8"))
    assert any("not password-protected" in n for n in report["notes"])
    # The journal never contains the password.
    for j in (tmp_path / "jobs").rglob("journal.json"):
        assert "secret-123" not in j.read_text(encoding="utf-8")


def test_encrypted_word_file(samples, tmp_path):
    res = batch([{"path": str(samples["docx_encrypted"]), "password": "secret-123"}], "txt", tmp_path / "o", tmp_path)
    r = res["results"][0]
    assert r["status"] == "done", r
    assert "Regional amounts" in Path(r["output"]).read_text(encoding="utf-8")


def test_outputs_are_never_overwritten(samples, tmp_path):
    out = tmp_path / "out"
    src = copy_sample(samples, "txt", tmp_path / "in")
    first = batch([{"path": str(src)}], "html", out, tmp_path, jobId="a")["results"][0]["output"]
    second = batch([{"path": str(src)}], "html", out, tmp_path, jobId="b")["results"][0]["output"]
    assert Path(first).name == "notes.html" and Path(second).name == "notes (1).html"
    assert Path(first).read_text(encoding="utf-8") and (out / "notes.html.report.html").exists()


@needs_lo
def test_cancel_then_resume_reuses_finished_work(samples, tmp_path, monkeypatch):
    """docx -> pptx (exact) runs three steps. Cancel during the second, resume, and only the rest runs again."""
    load_all()
    out = tmp_path / "out"
    src = copy_sample(samples, "docx", tmp_path / "in")
    calls: list[str] = []
    cancel = threading.Event()
    for eid in ("lo_to_pdf", "pdf_extract", "docmodel_pptx"):
        orig = STEPS[eid]

        def wrapped(ctx, art, target, _orig=orig, _eid=eid):
            calls.append(_eid)
            if _eid == "pdf_extract" and not cancel.is_set() and calls.count("pdf_extract") == 1:
                cancel.set()
                ctx.check = lambda: (_ for _ in ()).throw(__import__("otk_worker.errors", fromlist=["Cancelled"]).Cancelled())
                ctx.check()
            return _orig(ctx, art, target)
        monkeypatch.setitem(STEPS, eid, wrapped)
    res = batch([{"path": str(src)}], "pptx", out, tmp_path, cancel=cancel, jobId="resume-me", options={"mode": "exact"})
    assert res["results"][0]["status"] == "cancelled"
    assert res["jobDir"], "a cancelled job keeps its folder for resuming"
    assert calls == ["lo_to_pdf", "pdf_extract"]
    calls.clear()
    res = batch([{"path": str(src)}], "pptx", out, tmp_path, cancel=threading.Event(), jobId="resume-me",
                options={"mode": "exact"})
    assert res["results"][0]["status"] == "done", res["results"][0]
    assert calls == ["pdf_extract", "docmodel_pptx"], "LibreOffice's finished step was reused"
    assert res["jobDir"] is None, "a finished job removes its folder"


@needs_lo
def test_merge_each_target(samples, tmp_path):
    out = tmp_path / "out"
    src = {k: copy_sample(samples, k, tmp_path / k) for k in ("docx", "txt", "html")}
    pairs = {"pdf": ("docx", "txt"), "docx": ("html", "txt"), "html": ("docx", "txt"), "txt": ("docx", "html"),
             "epub": ("docx", "txt"), "png": ("docx", "txt"), "pptx": ("docx", "txt"), "xlsx": ("docx", "docx")}
    for target, (k1, k2) in pairs.items():
        files = [{"path": str(src[k1])}, {"path": str(src[k2] if k1 != k2 else copy_sample(samples, k2, tmp_path / "again"))}]
        res = batch(files, target, out / target, tmp_path, jobId=f"m-{target}", merge=True, mergeName="both")
        assert [r["status"] for r in res["results"]] == ["done", "done"], (target, res["results"])
        assert res["merged"] and res["merged"]["status"] == "done", (target, res["merged"])
        merged = Path(res["merged"]["output"])
        assert merged.exists(), target
        if target == "pdf":
            with pikepdf.open(merged) as pdf:
                with pdf.open_outline() as ol:
                    assert [i.title for i in ol.root] == [Path(r["output"]).stem for r in res["results"]]
        elif target == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(merged)
            assert len(wb.sheetnames) == 4 and len({n.lower() for n in wb.sheetnames}) == 4
            assert any("renamed" in n for n in res["merged"]["notes"])
        elif target == "pptx":
            from pptx import Presentation
            counts = [len(Presentation(r["output"]).slides) for r in res["results"]]
            assert len(Presentation(str(merged)).slides) == sum(counts)
        elif target == "png":
            assert merged.suffix == ".zip"


@needs_lo
@needs_gs
@needs_verapdf
@pytest.mark.parametrize("target", ["pdfa1b", "pdfa2b", "pdfa3b"])
def test_pdfa_flavours_are_proven_by_verapdf(samples, tmp_path, target):
    out = tmp_path / "out"
    files = [{"path": str(copy_sample(samples, k, tmp_path / k))} for k in ("docx", "pdf", "pdf_lo")]
    res = batch(files, target, out, tmp_path, jobId=target, options={"pdfaEmbedSource": True})
    for r in res["results"]:
        assert r["status"] == "done", r
        assert ".NOT-PDFA" not in r["output"], r
        v = pdfa.verapdf(Path(r["output"]), pdfa.FLAVOURS[target])
        assert v.compliant, (r["source"], v.failed_rules)
        if target == "pdfa3b":
            with pikepdf.open(r["output"]) as pdf:
                (name,) = list(pdf.attachments)
                assert pdf.attachments[name].obj.AFRelationship == "/Source"


@needs_tess
@needs_gs
@needs_verapdf
def test_scan_to_pdfa_gets_a_text_layer(samples, tmp_path):
    res = batch([{"path": str(samples["pdf_scanned"])}], "pdfa2b", tmp_path / "out", tmp_path, jobId="scan")
    r = res["results"][0]
    assert r["status"] == "done" and ".NOT-PDFA" not in r["output"], r
    import pypdfium2 as pdfium
    text = pdfium.PdfDocument(r["output"])[0].get_textpage().get_text_range()
    assert "Quarterly" in text and "001234" in text


@needs_lo
def test_failed_validation_is_never_called_pdfa(samples, tmp_path, monkeypatch):
    def reject(path, flavour, check=lambda: None):
        return pdfa.Validation(flavour, False, failed_rules=[{"clause": "6.1", "test": 1, "description": "test rule",
                                                             "failedChecks": 1, "specification": "ISO 19005-2"}])
    monkeypatch.setattr(pdfa, "verapdf", reject)
    res = batch([{"path": str(copy_sample(samples, "docx", tmp_path / "d"))}], "pdfa2b", tmp_path / "out", tmp_path)
    r = res["results"][0]
    assert r["status"] == "done" and r["output"].endswith("report.NOT-PDFA.pdf")
    assert r["verdict"] == "review"
    report = json.loads(Path(r["reportJson"]).read_text(encoding="utf-8"))
    assert next(c for c in report["checks"] if c["id"] == "pdfa")["status"] == "fail"

"""Converter building blocks: detection, HTML preparation, literal TXT, typing, PDF extraction, comparison rules."""

from __future__ import annotations

import base64
import zipfile
from pathlib import Path

import pytest
from conftest_converter import HAVE_LO, needs_lo

from otk_worker.converter import formats, htmlprep
from otk_worker.converter.context import Artifact, Options, StepContext
from otk_worker.converter.steps import load_all
from otk_worker.converter.steps.docmodel_out import typed_value
from otk_worker.converter.verify import compare
from otk_worker.errors import UnsupportedFormat

EXPECTED_FORMATS = {
    "docx": "docx", "xlsx": "xlsx", "pptx": "pptx", "html": "html", "txt": "txt", "svg": "svg", "pdf": "pdf",
    "png": "png", "jpeg": "jpeg", "pdf_scanned": "pdf_scanned", "epub": "epub",
    "doc": "doc", "xls": "xls", "ppt": "ppt", "pdfa2b": "pdfa2b", "pdf_lo": "pdf",
}


def test_detects_every_sample_by_content(samples, tmp_path):
    for key, fmt in EXPECTED_FORMATS.items():
        if key not in samples:
            assert not HAVE_LO, key
            continue
        assert formats.detect(samples[key]).format == fmt, key
    # The extension does not decide: a DOCX renamed to .pdf is still a DOCX.
    renamed = tmp_path / "really-word.pdf"
    renamed.write_bytes(samples["docx"].read_bytes())
    assert formats.detect(renamed).format == "docx"
    bogus = tmp_path / "archive.docx"
    with zipfile.ZipFile(bogus, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(UnsupportedFormat):
        formats.detect(bogus)


def test_detects_encryption_and_scans(samples):
    locked = formats.detect(samples["pdf_encrypted"])
    assert locked.encrypted and locked.pages is None
    opened = formats.detect(samples["pdf_encrypted"], "secret-123")
    assert opened.encrypted and opened.pages == 2
    assert formats.detect(samples["docx_encrypted"]).encrypted
    scan = formats.detect(samples["pdf_scanned"])
    assert scan.format == "pdf_scanned" and scan.scanned_pages == [0]


def test_html_preparation_inlines_local_and_blocks_remote(tmp_path):
    (tmp_path / "img").mkdir()
    png = base64.b64decode(htmlprep.PLACEHOLDER_PNG.split(",")[1])
    (tmp_path / "img" / "local pic.png").write_bytes(png)
    (tmp_path / "style.css").write_text("@import url('more.css'); body{background:url(img/local%20pic.png)}", encoding="utf-8")
    (tmp_path / "more.css").write_text("p{color:red} h1{background:url(https://cdn.example.com/bg.png)}", encoding="utf-8")
    src = tmp_path / "page.html"
    src.write_text("""<html><head><link rel="stylesheet" href="style.css">
<link rel="stylesheet" href="https://cdn.example.com/x.css"><script src="https://cdn.example.com/a.js"></script>
<script>alert(1)</script></head><body><img src="img/local%20pic.png"><img src="https://example.com/r.png" alt="r">
<img src="missing.png"><a href="https://example.com/link">a link stays a link</a></body></html>""", encoding="utf-8")
    out = tmp_path / "out.html"
    prep = htmlprep.prepare_html(src, out)
    text = out.read_text(encoding="utf-8")
    assert "cdn.example.com" not in text and "https://example.com/r.png" not in text.replace('data-otk-removed="https://example.com/r.png"', "")
    assert 'href="https://example.com/link"' in text          # links are not resources
    assert text.count("data:image/png;base64") >= 2           # local picture inlined (img and CSS)
    assert "<script" not in text and prep.scripts_removed == 2
    assert set(prep.blocked) >= {"https://cdn.example.com/x.css", "https://cdn.example.com/a.js",
                                 "https://example.com/r.png", "https://cdn.example.com/bg.png"}
    assert prep.missing == ["missing.png"]
    assert any("remote resource" in n for n in prep.notes())


def _ctx(tmp_path: Path, source: Path, fmt: str, **opts) -> StepContext:
    return StepContext(work=tmp_path / "work", options=Options.from_dict(opts), source=source, source_format=fmt)


def test_txt_is_literal_never_markdown(tmp_path):
    src = tmp_path / "notes.txt"
    lines = ["# not a heading", "*not bold* and _not italic_", "a\tb\tc", "नमस्ते दुनिया", "مرحبا بالعالم", "", "end"]
    src.write_text("\n".join(lines) + "\n", encoding="utf-8")
    steps = load_all()
    ctx = _ctx(tmp_path, src, "txt")
    html = steps["txt_html"](ctx, Artifact("txt", [src]), "html").path.read_text(encoding="utf-8")
    for ln in lines[:2]:
        assert ln in html
    assert 'dir="auto"' in html
    from docx import Document
    docx = steps["txt_docx"](ctx, Artifact("txt", [src]), "docx").path
    paras = [p.text for p in Document(str(docx)).paragraphs]
    assert paras == lines
    from openpyxl import load_workbook
    xlsx = steps["txt_xlsx"](ctx, Artifact("txt", [src]), "xlsx").path
    ws = load_workbook(xlsx).active
    assert [ws.cell(row=i + 1, column=1).value for i in range(len(lines))] == [ln or None for ln in lines]


def test_number_typing_is_conservative():
    assert typed_value("001234") == ("001234", "@")
    assert typed_value("3.50") == (3.5, "0.00")
    assert typed_value("12") == (12, "0")
    assert typed_value("1,234")[0] == "1,234"
    assert typed_value("2026-09-30")[0] == "2026-09-30"
    assert typed_value(" 7")[0] == " 7"


@needs_lo
def test_document_model_from_libreoffice_pdf(samples, tmp_path):
    from otk_worker.converter import docmodel

    m = docmodel.extract(samples["pdf_lo"], tmp_path / "m")
    p = m.pages[0]
    texts = [b.text for b in p.blocks if b.table is None]
    # Spaces LibreOffice draws in another font are put back between the Hindi words.
    assert "नमस्ते दुनिया यह हिंदी पाठ है" in texts
    assert any(b.lines[0].rtl for b in p.blocks if "مرحبا" in b.text)
    assert "A picture of the office:" in texts       # the paragraph after the table is not swallowed by it
    (t,) = p.tables
    assert (t.rows, t.cols) == (4, 3)
    cells = {(c.row, c.col): (c.rowspan, c.colspan, c.text) for c in t.cells}
    assert cells[(0, 0)] == (1, 3, "Regional amounts")
    assert cells[(2, 1)][2] == "001234" and cells[(2, 2)][2] == "3.50"
    linked = [s for b in p.blocks for ln in b.lines for s in ln.spans if s.link]
    assert [s.text for s in linked] == ["https://example.com/report"]   # only the linked characters
    assert len(p.images) == 1 and p.images[0].original


def test_text_comparison_rules():
    n = compare.normalize
    assert n("exam-\nple") == "exam-ple"                 # line-end hyphen kept, line joined
    assert n("soft­hyphen") == "softhyphen"
    assert n("ﬁne") == "fine"                             # NFKC
    assert "‍" in n("क्‍ष")                     # ZWJ kept (it matters for Indic shaping)
    assert n("हिंंदी") == "हिंदी"                          # a doubled combining mark is a reader artefact
    rtl = compare._canonical(["a", "مرحبا", "بالعالم", "b"])
    assert rtl == compare._canonical(["a", "بالعالم", "مرحبا", "b"])


def _ex(text, **kw):
    from otk_worker.converter.verify.extract import Extract
    return Extract("x", text=text, **kw)


def _sit(**kw):
    base = dict(source_format="docx", target="pdf", mode="exact", lost={}, target_caps={"text": "full"},
                expected_checks={}, added_text=[], expected_notes=[])
    base.update(kw)
    return compare.Situation(**base)


def test_text_check_verdicts():
    src = _ex("Total 15.75 in Q3")
    assert compare.check_text(src, _ex("Total 15.75 in Q3"), _sit()).status == "pass"
    missing = compare.check_text(src, _ex("Total in Q3"), _sit())
    assert missing.status == "fail" and missing.details["missing"][0]["token"] == "15.75"
    # Bullets, table rules and added labels are not content changes.
    assert compare.check_text(src, _ex("• Total | 15.75 | in Q3 +-----+"), _sit()).status == "pass"
    assert compare.check_text(src, _ex("Sheet: Sales Total 15.75 in Q3"), _sit(added_text=["Sheet: Sales"])).status == "pass"
    # Reordered words: a failure unless the route says the order changes.
    assert compare.check_text(src, _ex("in Q3 Total 15.75"), _sit()).status == "fail"
    assert compare.check_text(src, _ex("in Q3 Total 15.75"), _sit(target="xlsx")).status == "expected"
    # "Perfect" needs every check to pass and no declared change.
    ok = compare.Check("text", "Text", "pass", "")
    assert compare.verdict([ok], []) == "perfect"
    assert compare.verdict([ok], ["JPEG is lossy"]) == "expected"
    assert compare.verdict([ok, compare.Check("links", "Links", "fail", "")], []) == "review"


def test_engine_status_reports_bundled_versions(tmp_path, monkeypatch):
    import json

    from otk_worker.converter import engines

    (tmp_path / "pandoc").mkdir()
    (tmp_path / "pandoc" / "pandoc.exe").write_bytes(b"")
    (tmp_path / "pandoc" / "pandoc").write_bytes(b"")
    (tmp_path / "manifest.json").write_text(json.dumps({"pandoc": "pandoc 3.8.2.1", "ghostscript": "Ghostscript 10"}))
    monkeypatch.setenv("OTK_ENGINES", str(tmp_path))
    monkeypatch.setattr(engines, "_cache", {})
    st = engines.status()
    assert st["pandoc"]["bundled"] and st["pandoc"]["version"] == "pandoc 3.8.2.1"
    # Only bundled engines carry the bundle's version; one found on PATH is whatever is installed there.
    assert "version" not in st["gs"]


def _page_with_lines(drift: float = 0.0, drop_word: bool = False, move_line: int = 0):
    """A 100-DPI 'page' of text lines drawn word by word; `drift` px is added gradually along each line."""
    from PIL import Image, ImageDraw, ImageFont

    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 15)
    except OSError:
        font = ImageFont.load_default(15)
    img = Image.new("RGB", (850, 1100), "white")
    d = ImageDraw.Draw(img)
    words = "The quarterly figures show steady growth across every region with the strongest gains".split()
    for row in range(20):
        x, y = 60.0, 80 + row * 22 + (move_line if row == 7 else 0)
        for i, wd in enumerate(words):
            if not (drop_word and row == 7 and i == 5):
                d.text((round(x + drift * x / 850), y), wd, fill="black", font=font)
            x += d.textlength(wd + " ", font=font)
    return img


def test_appearance_tolerates_glyph_drift_but_not_changes():
    from otk_worker.common.ssim import ssim

    ref = _page_with_lines()
    score, _, regions = compare._compare_page(ref, _page_with_lines(drift=2.0), ssim)
    assert score >= 0.98 and not regions, "up to 2 px of drift along a line is how renderers differ, not a change"
    score, _, regions = compare._compare_page(ref, _page_with_lines(drop_word=True), ssim)
    assert regions, "a missing word must be found even though SSIM barely moves"
    x, y, w, h, _ = regions[0]
    assert 70 <= y + h / 2 <= 270 and w >= 30, regions[0]
    _, _, regions = compare._compare_page(ref, _page_with_lines(move_line=6), ssim)
    assert regions, "a line moved by 6 px must be found"
    score, _, regions = compare._compare_page(ref, _page_with_lines(drift=8.0), ssim)
    assert regions or score < 0.98, "drift beyond the 2 px tolerance must be found"


def test_appearance_hairlines_lighter_is_fine_missing_is_not():
    from PIL import Image, ImageDraw

    from otk_worker.common.ssim import ssim

    def page(shade):
        img = Image.new("RGB", (850, 1100), "white")
        if shade is not None:
            ImageDraw.Draw(img).line([(100, 400), (500, 400)], fill=(shade,) * 3, width=1)
        return img

    _, _, regions = compare._compare_page(page(0), page(110), ssim)
    assert not regions, "a rule drawn lighter (another resolution) is the same rule"
    _, _, regions = compare._compare_page(page(0), page(None), ssim)
    assert regions and regions[0][2] > 300, "a missing table rule must be found"


def test_missing_web_font_is_the_pages_own_fallback_but_missing_document_font_fails():
    from otk_worker.converter.verify.extract import Extract

    def sit(src_fmt):
        return compare.Situation(src_fmt, "pdf", "exact", {}, {}, {}, [], [])

    out = Extract("pdf", fonts_used={"DejaVu Sans": True})
    html = Extract("html", fonts_requested={"No Such Font Family 7"})
    c = compare.check_fonts(html, out, sit("html"))
    assert c.status == compare.EXPECTED and "fallback" in c.summary, c.summary
    docx = Extract("docx", fonts_requested={"No Such Font Family 7"})
    c = compare.check_fonts(docx, out, sit("docx"))
    assert c.status == compare.FAIL, c.summary

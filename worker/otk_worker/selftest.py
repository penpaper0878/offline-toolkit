"""Full offline self-test (Settings → Check every module).

One small real job in each module, from inputs made here (and the bundled sample portrait), with the
network guard on, so a pass shows that every engine, model and font the app needs is present and works
without the internet. The main process checks separately that no network attempt was made meanwhile.
"""

from __future__ import annotations

import json
import shutil
import threading
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import numpy as np
from PIL import Image, ImageDraw

from .errors import Cancelled


@dataclass
class Check:
    module: str
    name: str
    run: Callable[["Env"], str]


class Env:
    def __init__(self, work: Path, resources: Path, host, cancel: threading.Event | None):
        self.work = work
        self.resources = resources
        self.samples = resources / "samples"
        self.host = host
        self.cancel = cancel

    def check(self) -> None:
        if self.cancel is not None and self.cancel.is_set():
            raise Cancelled()


class Failed(Exception):
    pass


def _expect(ok: bool, message: str) -> None:
    if not ok:
        raise Failed(message)


def _photo(w: int, h: int, seed: int = 7) -> Image.Image:
    """A photo-like picture: smooth colour fields, edges and grain (compresses like a photograph)."""
    rng = np.random.default_rng(seed)
    y, x = np.mgrid[0:h, 0:w].astype(np.float32)
    img = np.stack([120 + 80 * np.sin(x / 97 + y / 151), 110 + 70 * np.cos(y / 83), 140 + 60 * np.sin((x + y) / 211)], -1)
    img += rng.normal(0, 9, img.shape)
    pil = Image.fromarray(np.clip(img, 0, 255).astype(np.uint8))
    d = ImageDraw.Draw(pil)
    r = max(4, min(w, h) // 6)
    for _ in range(12):
        x0, y0 = int(rng.integers(0, max(1, w - r))), int(rng.integers(0, max(1, h - r)))
        d.ellipse((x0, y0, x0 + int(rng.integers(r // 2, r * 2)), y0 + int(rng.integers(r // 2, r * 2))),
                  fill=tuple(int(v) for v in rng.integers(0, 255, 3)))
    return pil


# ------------------------------------------------------------------ Module 1
def resizer_check(env: Env) -> str:
    from .common import imageio, metadata
    from .resizer import pipeline

    src = env.work / "photo.jpg"
    _photo(1600, 1200).save(src, quality=95)
    settings = {"width": 240, "height": 240, "unit": "px", "dpi": 200, "sizeMode": "exactPixels", "fitDpi": False,
                "lockAspect": True, "fit": "crop", "padColor": "#FFFFFF", "format": "jpeg", "quality": 90,
                "sizeRange": {"min": 20, "max": 50, "unit": "KB"}, "subsampling": "auto", "stripMetadata": True,
                "allowPadding": True, "allowQuantize": False, "webpLossless": "auto", "saveOutOfRange": False,
                "naming": "{name}"}
    result, data = pipeline.process(imageio.open_image(str(src)), None, settings, check=env.check)
    rb = metadata.readback(data)
    _expect((rb["width"], rb["height"]) == (240, 240), f"made {rb['width']} × {rb['height']} px")
    _expect(tuple(rb["jfifDensity"] or ()) == (200, 200), f"DPI in the file: {rb['jfifDensity']}")
    _expect(20 * 1024 <= len(data) <= 50 * 1024, f"{len(data) / 1024:.1f} KB, outside 20–50 KB")
    return f"240 × 240 px, 200 DPI, {len(data) / 1024:.1f} KB ({result['attempts']} encodes)"


# ------------------------------------------------------------------ Module 2
def _convert(env: Env, source: Path, target: str, **opts) -> dict:
    from .converter import runner
    from .converter.context import Options

    out = env.work / f"out-{source.stem}-{target}"
    out.mkdir(exist_ok=True)
    res = runner.convert_file(source, target, Options.from_dict(opts), env.work / f"job-{source.stem}-{target}", out,
                              host=env.host, check=env.check)
    d = res.to_dict()
    _expect(d["status"] == "done", d.get("message") or d["status"])
    _expect(d["verdict"] in ("perfect", "expected"), f"verdict {d['verdict']}: {d.get('message', '')}")
    return d


def _report_checks(d: dict) -> dict:
    data = json.loads(Path(d["reportJson"]).read_text(encoding="utf-8"))
    return {c["id"]: c for c in data.get("checks", [])}


def word_to_pdfa(env: Env) -> str:
    import docx

    src = env.work / "letter.docx"
    doc = docx.Document()
    doc.add_heading("Offline Toolkit self-test", 1)
    doc.add_paragraph("English, हिन्दी and العربية in one document.")
    t = doc.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            t.cell(r, c).text = f"R{r + 1}C{c + 1}"
    doc.save(src)
    d = _convert(env, src, "pdfa2b", mode="exact")
    pdfa = _report_checks(d).get("pdfa")
    _expect(pdfa is not None and pdfa["status"] == "pass", f"veraPDF: {pdfa and pdfa.get('detail')}")
    return f"LibreOffice → PDF/A-2b, veraPDF pass, verdict {d['verdict']}"


def ocr_check(env: Env) -> str:
    from .design import assets

    face = assets.face("Roboto", 400)
    _expect(face is not None, "font Roboto missing")
    img = Image.new("RGB", (1400, 260), "white")
    ImageDraw.Draw(img).text((40, 80), "Quick brown foxes jump offline 2026", font=face.pil(64), fill="black")
    src = env.work / "scan.png"
    img.save(src, dpi=(300, 300))
    d = _convert(env, src, "txt", mode="editable", ocr=True, ocrLanguages=["eng"])
    text = Path(d["output"]).read_text(encoding="utf-8").lower().split()
    want = "quick brown foxes jump offline 2026".split()
    found = sum(w in text for w in want)
    _expect(found >= len(want) - 1, f"read {' '.join(text)[:80]!r}")
    return f"Tesseract read {found} of {len(want)} words"


def html_to_epub(env: Env) -> str:
    src = env.work / "page.html"
    src.write_text("<!doctype html><html lang='en'><head><meta charset='utf-8'><title>Self-test</title></head><body>"
                   "<h1>Chapter one</h1><p>A paragraph with <a href='#x'>a link</a>.</p><ul><li>One</li><li>Two</li></ul>"
                   "</body></html>", encoding="utf-8")
    d = _convert(env, src, "epub", mode="editable")
    return f"Pandoc → EPUB, verdict {d['verdict']}"


def html_to_pdf(env: Env) -> str:
    _expect(env.host is not None, "the app's Chromium printer is not available here")
    src = env.work / "invoice.html"
    src.write_text("<!doctype html><html><head><meta charset='utf-8'><title>Invoice</title></head><body>"
                   "<h1>Invoice 42</h1><table border='1'><tr><td>Item</td><td>Price</td></tr><tr><td>Paper</td><td>3.50</td></tr></table>"
                   "</body></html>", encoding="utf-8")
    d = _convert(env, src, "pdf", mode="exact")
    return f"The app's Chromium → PDF, verdict {d['verdict']}"


def svg_to_png(env: Env) -> str:
    src = env.work / "badge.svg"
    src.write_text("<svg xmlns='http://www.w3.org/2000/svg' width='200' height='120' viewBox='0 0 200 120'>"
                   "<rect x='10' y='10' width='180' height='100' rx='16' fill='#2f6fde'/>"
                   "<circle cx='60' cy='60' r='30' fill='#ffd60a'/></svg>", encoding="utf-8")
    d = _convert(env, src, "png", mode="exact", dpi=96)
    with Image.open(d["outputs"][0] if d["outputs"] else d["output"]) as im:
        _expect(im.size[0] >= 200, f"{im.size}")
        return f"resvg → PNG {im.size[0]} × {im.size[1]} px"


# ------------------------------------------------------------------ Module 3
def design_check(env: Env) -> str:
    from .design import assets, pipeline

    face = assets.face("Montserrat", 700)
    _expect(face is not None, "font Montserrat missing")
    img = Image.new("RGB", (1000, 420), "#ffffff")
    d = ImageDraw.Draw(img)
    d.text((60, 70), "Summer Festival", font=face.pil(84), fill="#1d3557")
    d.rounded_rectangle((60, 260, 420, 350), radius=20, fill="#e63946")
    src = env.work / "poster.png"
    img.save(src, dpi=(150, 150))
    scene = pipeline.analyze(str(src), env.work / "design", langs=["eng"], upscale="never", check=env.check)
    texts = [ly for ly in scene["layers"] if ly["type"] == "text"]
    hit = next((t for t in texts if t.get("text", "").strip() == "Summer Festival"), None)
    _expect(hit is not None, f"text read: {[t.get('text') for t in texts]}")
    _expect(hit["style"]["family"] == "Montserrat", f"font {hit['style']['family']}")
    shapes = [ly for ly in scene["layers"] if ly["type"] == "shape"]
    _expect(bool(shapes), "the button was not found as a shape")
    return f"text read and matched to Montserrat {hit['style'].get('weight')}, {len(shapes)} shape(s), {len(scene['layers'])} layers"


def superres_check(env: Env) -> str:
    from .design import superres

    out, method = superres.upscale(_photo(96, 64), 2, check=env.check)
    _expect(out.size == (192, 128), f"{out.size}")
    _expect("esrgan" in method.lower(), f"method {method}")
    return f"Real-ESRGAN enlarged 96 × 64 to {out.size[0]} × {out.size[1]} px"


# ------------------------------------------------------------------ Module 4
def _portrait(env: Env) -> Path:
    p = env.samples / "passport" / "portrait.jpg"
    _expect(p.is_file(), f"sample portrait missing ({p})")
    return p


def segmentation_check(env: Env) -> str:
    from .design import cutout

    rgb = np.asarray(Image.open(_portrait(env)).convert("RGB"))
    prob = cutout.person_probability(rgb)
    _expect(prob is not None, "the person segmentation model is missing")
    share = float((prob > 0.5).mean())
    _expect(0.15 < share < 0.9, f"person covers {share:.0%} of the photo")
    return f"person found on {share:.0%} of the sample portrait"


def passport_check(env: Env) -> str:
    from .common import metadata
    from .passport import api as passport

    class Ctx:
        cancel_event = env.cancel or threading.Event()

    specs = json.loads((env.resources / "defaults" / "presets" / "passport-specs.json").read_text(encoding="utf-8"))["specs"]
    uk = next(s for s in specs if s["id"] == "uk-passport")
    o = passport.open_photo({"path": str(_portrait(env)), "previewDir": str(env.work)}, Ctx)
    out = env.work / "passport.jpg"
    res = passport.export_photo({"id": o["id"], "spec": uk, "background": {"mode": "replace", "color": "#FFFFFF", "feather": 1},
                                 "format": "jpeg", "path": str(out), "previewDir": str(env.work)}, Ctx)
    rb = metadata.readback(out.read_bytes())
    _expect((rb["width"], rb["height"]) == (413, 531), f"{rb['width']} × {rb['height']} px")
    _expect(tuple(rb["jfifDensity"] or ()) == (300, 300), f"DPI {rb['jfifDensity']}")
    hints = {h["id"]: h["level"] for h in res["hints"]}
    _expect(hints.get("face") == "ok" and hints.get("head") == "ok", f"hints {hints}")
    return "face found, head within the UK range, 413 × 531 px at 300 DPI"


CHECKS: list[Check] = [
    Check("resizer", "Resize to 240 × 240 px at 200 DPI within 20–50 KB", resizer_check),
    Check("converter", "Word to PDF/A-2b, validated", word_to_pdfa),
    Check("converter", "Read text from a picture (OCR)", ocr_check),
    Check("converter", "Web page to EPUB", html_to_epub),
    Check("converter", "Web page to PDF", html_to_pdf),
    Check("converter", "SVG to PNG", svg_to_png),
    Check("design", "Read a poster into layers", design_check),
    Check("design", "Enlarge a small picture (super-resolution)", superres_check),
    Check("passport", "Find the person in a photo", segmentation_check),
    Check("passport", "Make a UK passport photo", passport_check),
]


def run(params: dict, notify: Callable[[dict], None], cancel: threading.Event | None = None, host=None) -> dict:
    root = Path(params["workDir"])
    work = root / time.strftime("%Y%m%d-%H%M%S")
    shutil.rmtree(root, ignore_errors=True)
    work.mkdir(parents=True, exist_ok=True)
    env = Env(work, Path(params["resourcesDir"]), host, cancel)
    only = set(params.get("only") or [])
    checks = [c for c in CHECKS if not only or c.module in only]
    out = []
    try:
        for i, c in enumerate(checks):
            env.check()
            notify({"fraction": i / len(checks), "message": c.name})
            t0 = time.monotonic()
            try:
                detail, passed = c.run(env), True
            except Cancelled:
                raise
            except Exception as exc:  # noqa: BLE001 - every failure is reported, the next check still runs
                detail, passed = f"{type(exc).__name__ if not isinstance(exc, Failed) else 'Failed'}: {exc}", False
            out.append({"module": c.module, "name": c.name, "passed": passed, "detail": detail,
                        "seconds": round(time.monotonic() - t0, 1)})
        notify({"fraction": 1.0, "message": "Done"})
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return {"checks": out, "passed": all(c["passed"] for c in out) and bool(out)}

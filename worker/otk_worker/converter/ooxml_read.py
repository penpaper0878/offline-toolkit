"""Reading DOCX, PPTX and XLSX content in document order (for conversion steps and for verification).

The readers look at the XML directly where the python-docx/python-pptx APIs
skip content: text boxes, nested tables, group shapes. Word's
mc:AlternateContent Fallback branch (a VML copy of the same drawing) is
ignored so text is never counted twice.
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _q(ns: str, tag: str) -> str:
    return f"{{{ns}}}{tag}"


def _in_fallback(el) -> bool:
    p = el.getparent()
    while p is not None:
        if p.tag == _q(MC, "Fallback"):
            return True
        p = p.getparent()
    return False


# ------------------------------------------------------------------ DOCX
@dataclass
class DocxRun:
    text: str
    bold: bool = False
    italic: bool = False
    link: str | None = None


@dataclass
class DocxPara:
    runs: list[DocxRun]
    style: str = ""
    level: int = 0           # heading level (0 = body)
    rtl: bool = False
    list_level: int | None = None

    @property
    def text(self) -> str:
        return "".join(r.text for r in self.runs)


@dataclass
class Grid:
    rows: int
    cols: int
    cells: dict[tuple[int, int], str]               # top-left (row, col) -> text
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)   # r0, c0, r1, c1 (inclusive)
    links: dict[tuple[int, int], str] = field(default_factory=dict)       # first web link in a cell


def _para_text_runs(p, rels: dict[str, str]) -> list[DocxRun]:
    runs: list[DocxRun] = []
    for el in p.iter(_q(W, "r")):
        if _in_fallback(el):
            continue
        # Runs inside a text box belong to the text box's own paragraphs, not this one.
        anc = el.getparent()
        nested = False
        while anc is not None and anc is not p:
            if anc.tag in (_q(W, "txbxContent"), _q(W, "p")):
                nested = True
                break
            anc = anc.getparent()
        if nested:
            continue
        parts = []
        for ch in el:
            if ch.tag == _q(W, "t"):
                parts.append(ch.text or "")
            elif ch.tag == _q(W, "tab"):
                parts.append("\t")
            elif ch.tag in (_q(W, "br"), _q(W, "cr")):
                parts.append("\n")
            elif ch.tag == _q(W, "noBreakHyphen"):
                parts.append("‑")
            elif ch.tag == _q(W, "sym"):
                parts.append(chr(int(ch.get(_q(W, "char"), "20"), 16)))
        if not parts:
            continue
        rpr = el.find(_q(W, "rPr"))
        bold = rpr is not None and rpr.find(_q(W, "b")) is not None and rpr.find(_q(W, "b")).get(_q(W, "val"), "1") not in ("0", "false")
        italic = rpr is not None and rpr.find(_q(W, "i")) is not None and rpr.find(_q(W, "i")).get(_q(W, "val"), "1") not in ("0", "false")
        link = None
        h = el.getparent()
        if h is not None and h.tag == _q(W, "hyperlink"):
            rid = h.get(_q(R, "id"))
            link = rels.get(rid) if rid else (("#" + h.get(_q(W, "anchor"))) if h.get(_q(W, "anchor")) else None)
        runs.append(DocxRun("".join(parts), bool(bold), bool(italic), link))
    return runs


def _style_level(style_id: str, names: dict[str, str]) -> int:
    name = names.get(style_id, style_id).lower().replace(" ", "")
    if name == "title":
        return 1
    if name.startswith("heading") and name[7:].isdigit():
        return min(6, int(name[7:]))
    return 0


def docx_rels(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as z:
        try:
            root = etree.fromstring(z.read("word/_rels/document.xml.rels"))
        except KeyError:
            return {}
    return {r.get("Id"): r.get("Target") for r in root if r.get("TargetMode") == "External"}


def docx_style_names(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as z:
        try:
            root = etree.fromstring(z.read("word/styles.xml"))
        except KeyError:
            return {}
    out = {}
    for st in root.iter(_q(W, "style")):
        name = st.find(_q(W, "name"))
        out[st.get(_q(W, "styleId"))] = name.get(_q(W, "val")) if name is not None else ""
    return out


def _para(p, rels, names) -> DocxPara:
    ppr = p.find(_q(W, "pPr"))
    style, rtl, list_level = "", False, None
    if ppr is not None:
        ps = ppr.find(_q(W, "pStyle"))
        style = ps.get(_q(W, "val")) if ps is not None else ""
        rtl = ppr.find(_q(W, "bidi")) is not None
        num = ppr.find(_q(W, "numPr"))
        if num is not None:
            ilvl = num.find(_q(W, "ilvl"))
            list_level = int(ilvl.get(_q(W, "val"), "0")) if ilvl is not None else 0
    level = _style_level(style, names)
    if ppr is not None and not level and ppr.find(_q(W, "outlineLvl")) is not None:
        val = int(ppr.find(_q(W, "outlineLvl")).get(_q(W, "val"), "9"))
        level = val + 1 if val < 9 else 0
    return DocxPara(_para_text_runs(p, rels), style, level, rtl, list_level)


def table_grid(tbl, cell_text, cell_link=None) -> Grid:
    """Word table -> grid with merges (gridSpan horizontally, vMerge vertically)."""
    cells: dict[tuple[int, int], str] = {}
    links: dict[tuple[int, int], str] = {}
    merges = []
    open_v: dict[int, list] = {}   # col -> [r0, c0, r1, c1]
    r = 0
    max_c = 0
    for tr in tbl.findall(_q(W, "tr")):
        c = 0
        for tc in tr.findall(_q(W, "tc")):
            tcpr = tc.find(_q(W, "tcPr"))
            span = 1
            vmerge = None
            if tcpr is not None:
                gs = tcpr.find(_q(W, "gridSpan"))
                if gs is not None:
                    span = max(1, int(gs.get(_q(W, "val"), "1")))
                vm = tcpr.find(_q(W, "vMerge"))
                if vm is not None:
                    vmerge = vm.get(_q(W, "val"), "continue")
            if vmerge == "continue" and c in open_v:
                open_v[c][2] = r
            else:
                cells[(r, c)] = cell_text(tc)
                if cell_link is not None:
                    link = cell_link(tc)
                    if link:
                        links[(r, c)] = link
                if vmerge == "restart":
                    open_v[c] = [r, c, r, c + span - 1]
                else:
                    open_v.pop(c, None)
                    if span > 1:
                        merges.append((r, c, r, c + span - 1))
            c += span
        max_c = max(max_c, c)
        r += 1
    for m in open_v.values():
        if m[2] > m[0] or m[3] > m[1]:
            merges.append(tuple(m))
    # A vertical merge that also spans columns was recorded once; drop the duplicate horizontal one.
    merges = sorted(set(merges))
    clean = []
    for m in merges:
        if any(o != m and o[0] == m[0] and o[1] == m[1] and o[2] >= m[2] and o[3] >= m[3] for o in merges):
            continue
        clean.append(m)
    return Grid(r, max_c, cells, clean, links)


def docx_body(path: Path) -> list[tuple[str, object]]:
    """[("p", DocxPara) | ("table", Grid)] in body order; text-box paragraphs follow their anchor paragraph."""
    rels = docx_rels(path)
    names = docx_style_names(path)
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(_q(W, "body"))
    out: list[tuple[str, object]] = []

    def cell_text(tc) -> str:
        parts = []
        for p in tc.iter(_q(W, "p")):
            if _in_fallback(p):
                continue
            parts.append(_para(p, rels, names).text)
        return "\n".join(parts).strip("\n")

    def cell_link(tc) -> str | None:
        for p in tc.iter(_q(W, "p")):
            for run in _para(p, rels, names).runs:
                if run.link and not run.link.startswith("#"):
                    return run.link
        return None

    def walk(container):
        for el in container:
            if el.tag == _q(W, "p"):
                out.append(("p", _para(el, rels, names)))
                for tx in el.iter(_q(W, "txbxContent")):
                    if _in_fallback(tx):
                        continue
                    for p in tx.findall(_q(W, "p")):
                        out.append(("p", _para(p, rels, names)))
            elif el.tag == _q(W, "tbl"):
                out.append(("table", table_grid(el, cell_text, cell_link)))
            elif el.tag == _q(W, "sdt"):
                content = el.find(_q(W, "sdtContent"))
                if content is not None:
                    walk(content)
    walk(body)
    return out


# ------------------------------------------------------------------ PPTX
@dataclass
class SlideItem:
    kind: str                          # text | table | picture
    paragraphs: list[DocxPara] = field(default_factory=list)
    grid: Grid | None = None
    image: bytes | None = None
    image_ext: str = "png"
    size_pt: tuple[float, float] = (0.0, 0.0)
    is_title: bool = False


@dataclass
class SlideContent:
    index: int
    items: list[SlideItem]
    notes: list[str]
    title: str = ""


def _pptx_paragraphs(text_frame) -> list[DocxPara]:
    out = []
    for p in text_frame.paragraphs:
        runs = []
        for el in p._p:
            tag = etree.QName(el).localname
            if tag == "r":
                r = next((x for x in p.runs if x._r is el), None)
                if r is None:
                    continue
                link = None
                try:
                    link = r.hyperlink.address
                except Exception:
                    pass
                runs.append(DocxRun(r.text, bool(r.font.bold), bool(r.font.italic), link))
            elif tag == "br":
                runs.append(DocxRun("\n"))
            elif tag == "fld":
                t = el.find(_q(A, "t"))
                runs.append(DocxRun(t.text if t is not None and t.text else ""))
        ppr = p._p.pPr
        rtl = ppr is not None and ppr.get("rtl") == "1"
        out.append(DocxPara(runs, level=0, rtl=rtl, list_level=p.level))
    return out


def _pptx_grid(table) -> Grid:
    cells, merges, links = {}, [], {}
    rows, cols = len(table.rows), len(table.columns)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            if cell.is_spanned:
                continue
            paras = _pptx_paragraphs(cell.text_frame)
            cells[(r, c)] = "\n".join(p.text for p in paras).strip("\n")
            link = next((x.link for p in paras for x in p.runs if x.link and not x.link.startswith("#")), None)
            if link:
                links[(r, c)] = link
            if cell.is_merge_origin:
                merges.append((r, c, r + cell.span_height - 1, c + cell.span_width - 1))
    return Grid(rows, cols, cells, merges, links)


def pptx_slides(path: Path) -> list[SlideContent]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        items: list[tuple[float, float, SlideItem]] = []

        def visit(shapes, dx=0, dy=0):
            for sh in shapes:
                top = (sh.top or 0) + dy
                left = (sh.left or 0) + dx
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    visit(sh.shapes, dx, dy)
                    continue
                is_title = False
                if sh.is_placeholder:
                    try:
                        is_title = sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
                    except Exception:
                        pass
                if getattr(sh, "has_table", False) and sh.has_table:
                    items.append((top, left, SlideItem("table", grid=_pptx_grid(sh.table))))
                elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE or (hasattr(sh, "image") and _has_image(sh)):
                    try:
                        img = sh.image
                        items.append((top, left, SlideItem("picture", image=img.blob, image_ext=img.ext,
                                                           size_pt=((sh.width or 0) / 12700, (sh.height or 0) / 12700))))
                    except Exception:
                        pass
                elif getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
                    items.append((top, left, SlideItem("text", paragraphs=_pptx_paragraphs(sh.text_frame),
                                                       is_title=is_title)))
        visit(slide.shapes)
        items.sort(key=lambda t: (not t[2].is_title, round(t[0] / 12700 / 4), t[1]))
        notes = []
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if tf is not None:
                notes = [p.text for p in _pptx_paragraphs(tf)]
                while notes and not notes[-1].strip():
                    notes.pop()
        title = next((("".join(p.text for p in it.paragraphs)).strip() for _, _, it in items if it.is_title), "")
        slides.append(SlideContent(i, [it for _, _, it in items], notes, title))
    return slides


def _has_image(sh) -> bool:
    try:
        sh.image
        return True
    except Exception:
        return False

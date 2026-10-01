"""Merging the converted files of a batch into one file of the target format.

PDF/PDF-A: pages in order, one top-level bookmark per file with its own bookmarks nested under it
(PDF/A is re-declared and re-validated). DOCX: docxcompose, a page break between files. XLSX: sheets
copied (values, formulas, styles, merged cells, widths); a clashing sheet name is renamed and reported.
PPTX: slides copied (shapes, pictures, speaker notes). HTML/TXT/EPUB: joined in order. Images and SVG:
one ZIP.
"""

from __future__ import annotations

import copy
import html as htmlmod
import os
import shutil
import zipfile
from pathlib import Path
from typing import Callable

from lxml import html as lhtml

from ..common.paths import safe_stem, unique_path
from ..errors import ToolkitError
from . import formats, pdfa


def merge_outputs(paths: list[Path], target: str, out_dir: Path, name: str, *, work: Path,
                  check: Callable[[], None] = lambda: None) -> dict:
    work.mkdir(parents=True, exist_ok=True)
    stem = safe_stem(name, "merged")
    ext = formats.OUTPUT_EXT[target]
    notes: list[str] = []
    zipped = any(p.suffix.lower() == ".zip" for p in paths)
    if target in ("png", "jpeg", "svg") or zipped:
        out = unique_path(out_dir, stem, ".zip")
        _zip_all(paths, out)
        return {"status": "done", "output": str(out), "notes": ["The files are collected in one ZIP."]}
    tmp = work / f"merged{ext}"
    if target in ("pdf", "pdfa1b", "pdfa2b", "pdfa3b"):
        _pdf(paths, tmp)
        if target in pdfa.FLAVOURS:
            flavour = pdfa.FLAVOURS[target]
            fixed = work / "merged-pdfa.pdf"
            pdfa.finish(tmp, fixed, flavour)
            v = pdfa.verapdf(fixed, flavour, check=check)
            tmp = fixed
            if not v.compliant:
                stem += ".NOT-PDFA"
                notes.append(f"The merged file did not pass veraPDF ({len(v.failed_rules)} rule(s)); it is saved as NOT-PDFA.")
            else:
                notes.append(f"The merged file passed veraPDF (PDF/A-{flavour}).")
    elif target == "docx":
        _docx(paths, tmp)
    elif target == "xlsx":
        notes += _xlsx(paths, tmp)
    elif target == "pptx":
        notes += _pptx(paths, tmp)
    elif target == "html":
        _html(paths, tmp, name)
    elif target == "txt":
        tmp.write_text("\n\f\n".join(p.read_text(encoding="utf-8").rstrip("\n") for p in paths) + "\n", encoding="utf-8")
        notes.append("Files are separated by a form feed (page break) character.")
    elif target == "epub":
        _epub(paths, tmp, name)
    else:
        raise ToolkitError(f"Merging is not available for {target.upper()}.")
    out = unique_path(out_dir, stem, ext)
    shutil.copyfile(tmp, out)
    return {"status": "done", "output": str(out), "notes": notes, "inputs": [str(p) for p in paths]}


def _zip_all(paths: list[Path], out: Path) -> None:
    part = out.with_name(out.name + ".part")
    with zipfile.ZipFile(part, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            if p.suffix.lower() == ".zip":
                with zipfile.ZipFile(p) as src:
                    for info in src.infolist():
                        z.writestr(f"{p.stem}/{info.filename}", src.read(info))
            else:
                z.write(p, arcname=p.name)
    os.replace(part, out)


def _pdf(paths: list[Path], out: Path) -> None:
    import pikepdf
    from pikepdf import OutlineItem

    merged = pikepdf.new()
    tops = []
    for p in paths:
        with pikepdf.open(p) as src:
            offset = len(merged.pages)
            children = []
            with src.open_outline() as ol:
                children = [_copy_item(it, src, offset) for it in ol.root]
            merged.pages.extend(src.pages)
            top = OutlineItem(p.stem, offset)
            top.children.extend(c for c in children if c is not None)
            tops.append(top)
    with merged.open_outline() as ol:
        ol.root.extend(tops)
    merged.save(out)


def _copy_item(item, src, offset: int):
    from pikepdf import OutlineItem

    page = None
    try:
        dest = item.destination
        if dest is not None and len(dest) > 0:
            target = dest[0]
            for i, pg in enumerate(src.pages):
                if pg.obj.objgen == target.objgen:
                    page = i
                    break
    except Exception:
        page = None
    new = OutlineItem(item.title, (page if page is not None else 0) + offset)
    for ch in item.children:
        c = _copy_item(ch, src, offset)
        if c is not None:
            new.children.append(c)
    return new


def _docx(paths: list[Path], out: Path) -> None:
    from docx import Document
    from docxcompose.composer import Composer

    master = Document(str(paths[0]))
    composer = Composer(master)
    for p in paths[1:]:
        doc = Document(str(p))
        if doc.paragraphs:
            doc.paragraphs[0].paragraph_format.page_break_before = True
        composer.append(doc)
    composer.save(str(out))


def _xlsx(paths: list[Path], out: Path) -> list[str]:
    from openpyxl import Workbook, load_workbook

    notes = []
    dest = Workbook()
    dest.remove(dest.active)
    used: set[str] = set()
    for p in paths:
        wb = load_workbook(p)
        for ws in wb.worksheets:
            title = ws.title
            n = 2
            while title.lower() in used:
                suffix = f" ({n})"
                title = ws.title[: 31 - len(suffix)] + suffix
                n += 1
            if title != ws.title:
                notes.append(f"Sheet '{ws.title}' from {p.name} was renamed to '{title}' (name already used).")
            used.add(title.lower())
            new = dest.create_sheet(title)
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None and not c.has_style:
                        continue
                    nc = new.cell(row=c.row, column=c.column, value=c.value)
                    if c.has_style:
                        nc.font = copy.copy(c.font)
                        nc.fill = copy.copy(c.fill)
                        nc.border = copy.copy(c.border)
                        nc.alignment = copy.copy(c.alignment)
                        nc.protection = copy.copy(c.protection)
                        nc.number_format = c.number_format
                    if c.hyperlink is not None:
                        nc.hyperlink = copy.copy(c.hyperlink)
            for rng in ws.merged_cells.ranges:
                new.merge_cells(str(rng))
            for key, dim in ws.column_dimensions.items():
                new.column_dimensions[key].width = dim.width
            for key, dim in ws.row_dimensions.items():
                new.row_dimensions[key].height = dim.height
            if ws.sheet_state != "visible":
                new.sheet_state = ws.sheet_state
            if getattr(ws, "_images", None):
                notes.append(f"Pictures on sheet '{ws.title}' of {p.name} are not copied into the merged workbook.")
    notes.append("Formulas that refer to other sheets keep their sheet names; check them if a sheet was renamed.")
    dest.save(out)
    return notes


def _pptx(paths: list[Path], out: Path) -> list[str]:
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    notes = []
    dest = Presentation(str(paths[0]))
    blank = dest.slide_layouts[6] if len(dest.slide_layouts) > 6 else dest.slide_layouts[-1]
    for p in paths[1:]:
        src = Presentation(str(p))
        if (src.slide_width, src.slide_height) != (dest.slide_width, dest.slide_height):
            notes.append(f"{p.name} has a different slide size; its slides keep their positions on the first file's size.")
        for slide in src.slides:
            new = dest.slides.add_slide(blank)
            for shp in list(new.shapes):
                shp._element.getparent().remove(shp._element)
            rid_map = {}
            for rel in slide.part.rels.values():
                if rel.reltype in (RT.IMAGE, RT.MEDIA, RT.VIDEO) and not rel.is_external:
                    rid_map[rel.rId] = new.part.relate_to(rel.target_part, rel.reltype)
                elif rel.reltype == RT.HYPERLINK and rel.is_external:
                    rid_map[rel.rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            for el in slide.shapes._spTree.iterchildren():
                tag = el.tag.split("}")[-1]
                if tag in ("nvGrpSpPr", "grpSpPr"):
                    continue
                new_el = copy.deepcopy(el)
                for node in new_el.iter():
                    for attr, val in list(node.attrib.items()):
                        if attr.endswith("}embed") or attr.endswith("}link") or attr.endswith("}id"):
                            if val in rid_map:
                                node.set(attr, rid_map[val])
                new.shapes._spTree.append(new_el)
            bg = slide._element.cSld.bg
            if bg is not None:
                new._element.cSld.insert(0, copy.deepcopy(bg))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                new.notes_slide.notes_text_frame.text = slide.notes_slide.notes_text_frame.text
        notes.append(f"Slides from {p.name} use the first file's theme and layouts; charts and embedded "
                     "objects other than pictures are not copied.")
    dest.save(out)
    return notes


def _html(paths: list[Path], out: Path, name: str) -> None:
    styles, sections = [], []
    for i, p in enumerate(paths, start=1):
        doc = lhtml.document_fromstring(p.read_text(encoding="utf-8"))
        for st in doc.iter("style"):
            if st.text:
                styles.append(st.text)
        body = doc.find("body")
        inner = "".join(lhtml.tostring(ch, encoding="unicode") for ch in (body if body is not None else []))
        sections.append(f'<section class="merged-file" id="file-{i}" data-source="{htmlmod.escape(p.name, quote=True)}">'
                        f"{inner}</section>")
    out.write_text(f"<!DOCTYPE html>\n<html><head><meta charset=\"utf-8\"><title>{htmlmod.escape(name)}</title>"
                   f"<style>{''.join(styles)} section.merged-file{{break-before:page}}</style></head><body>\n"
                   + "\n<hr>\n".join(sections) + "\n</body></html>\n", encoding="utf-8")


def _epub(paths: list[Path], out: Path, name: str) -> None:
    import posixpath
    import uuid
    from datetime import datetime, timezone

    from lxml import etree

    from .epub import DC, OPF, open_book

    manifest, spine, nav_items, files = [], [], [], []
    fixed_any = False
    for b, p in enumerate(paths, start=1):
        book = open_book(p)
        zf = book.zf
        container = etree.fromstring(zf.read("META-INF/container.xml"))
        rootfile = container.find(".//{urn:oasis:names:tc:opendocument:xmlns:container}rootfile").get("full-path")
        opf = etree.fromstring(zf.read(rootfile))
        fixed = b"pre-paginated" in zf.read(rootfile)
        fixed_any = fixed_any or fixed
        prefix = f"book{b}"
        for info in zf.infolist():
            if info.filename.startswith(book.opf_dir + "/" if book.opf_dir else "") and not info.filename.endswith(".opf"):
                rel = info.filename[len(book.opf_dir) + 1:] if book.opf_dir else info.filename
                if info.filename.startswith(("META-INF/", "mimetype")):
                    continue
                files.append((f"{prefix}/{rel}", zf.read(info)))
        items = {}
        for it in opf.iter(f"{{{OPF}}}item"):
            props = " ".join(x for x in (it.get("properties") or "").split() if x != "nav")
            iid = f"{prefix}-{it.get('id')}"
            items[it.get("id")] = iid
            manifest.append(f'<item id="{iid}" href="{prefix}/{it.get("href")}" media-type="{it.get("media-type")}"'
                            + (f' properties="{props}"' if props else "") + "/>")
        for ref in opf.iter(f"{{{OPF}}}itemref"):
            layout = ' properties="rendition:layout-pre-paginated"' if fixed else ""
            spine.append(f'<itemref idref="{items[ref.get("idref")]}"{layout}/>')
        first = next((it for it in opf.iter(f"{{{OPF}}}itemref")), None)
        href = None
        if first is not None:
            item = next(it for it in opf.iter(f"{{{OPF}}}item") if it.get("id") == first.get("idref"))
            href = f"{prefix}/{item.get('href')}"
        sub = "".join(f'<li><a href="{prefix}/{posixpath.normpath(t[1])}">{htmlmod.escape(t[0])}</a></li>'
                      for t in book.nav if t[0] and t[1])
        nav_items.append(f'<li><a href="{href or "#"}">{htmlmod.escape(book.title or p.stem)}</a>'
                         + (f"<ol>{sub}</ol>" if sub else "") + "</li>")
    nav = (f'<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n<html xmlns="http://www.w3.org/1999/xhtml" '
           f'xmlns:epub="http://www.idpf.org/2007/ops"><head><meta charset="utf-8"/><title>Contents</title></head><body>'
           f'<nav epub:type="toc" id="toc"><ol>{"".join(nav_items)}</ol></nav></body></html>')
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    opf_xml = (f'<?xml version="1.0" encoding="utf-8"?>\n<package xmlns="{OPF}" version="3.0" unique-identifier="bookid" '
               f'prefix="rendition: http://www.idpf.org/vocab/rendition/#"><metadata xmlns:dc="{DC}">'
               f'<dc:identifier id="bookid">urn:uuid:{uuid.uuid4()}</dc:identifier><dc:title>{htmlmod.escape(name)}</dc:title>'
               f'<dc:language>en</dc:language><meta property="dcterms:modified">{modified}</meta></metadata>'
               f'<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
               f'{"".join(manifest)}</manifest><spine>{"".join(spine)}</spine></package>')
    container = ('<?xml version="1.0" encoding="utf-8"?><container version="1.0" '
                 'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                 '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    with zipfile.ZipFile(out, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", container, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/content.opf", opf_xml, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/nav.xhtml", nav, compress_type=zipfile.ZIP_DEFLATED)
        for rel, data in files:
            z.writestr(f"OEBPS/{rel}", data, compress_type=zipfile.ZIP_DEFLATED)

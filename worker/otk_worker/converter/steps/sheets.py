"""Spreadsheet steps (XLSX to HTML/TXT/DOCX/PPTX) and Word tables to sheets.

Displayed values come from LibreOffice ("save cell contents as shown": number
formats applied, formulas recalculated), so 3.5 formatted as 0.00 stays "3.50".
Formulas, number formats and raw values come from openpyxl and travel in
data- attributes of the semantic HTML, so HTML -> XLSX can restore them.
"""

from __future__ import annotations

import csv
import html
from dataclasses import dataclass
from pathlib import Path

from .. import libreoffice as lo
from .. import ooxml, ooxml_read
from ..context import Artifact, StepContext
from . import step
from .docmodel_out import safe_sheet_title, set_link, write_cell


@dataclass
class SheetData:
    name: str
    shown: list[list[str]]                       # displayed values, row-major
    formulas: dict[tuple[int, int], str]
    formats: dict[tuple[int, int], str]
    values: dict[tuple[int, int], object]
    merges: list[tuple[int, int, int, int]]      # 0-based inclusive
    hidden: bool = False

    @property
    def rows(self) -> int:
        return len(self.shown)

    @property
    def cols(self) -> int:
        return max((len(r) for r in self.shown), default=0)


def read_workbook(ctx: StepContext, path: Path) -> list[SheetData]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False)
    names = wb.sheetnames
    shown_files = lo.sheet_csvs(path, ctx.folder("csv"), names, check=ctx.check)
    sheets = []
    for ws in wb.worksheets:
        ctx.check()
        shown: list[list[str]] = []
        f = shown_files.get(ws.title)
        if f is not None:
            with open(f, encoding="utf-8", newline="") as fh:
                shown = [list(r) for r in csv.reader(fh, delimiter="\t", quotechar='"')]
        formulas, formats, values = {}, {}, {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                key = (c.row - 1, c.column - 1)
                if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
                    formulas[key] = str(c.value)
                else:
                    values[key] = c.value
                if c.number_format and c.number_format != "General":
                    formats[key] = c.number_format
        rows = len(shown)
        cols = max((len(r) for r in shown), default=0)
        if values or formulas:
            rows = max(rows, max(k[0] for k in [*values, *formulas]) + 1)
            cols = max(cols, max(k[1] for k in [*values, *formulas]) + 1)
        grid = [[(shown[r][c] if r < len(shown) and c < len(shown[r]) else "") for c in range(cols)] for r in range(rows)]
        # Trim empty trailing rows/columns (formatting-only cells).
        while grid and not any(x.strip() for x in grid[-1]):
            grid.pop()
        width = max((max((i + 1 for i, x in enumerate(r) if x != ""), default=0) for r in grid), default=0)
        grid = [r[:width] + [""] * (width - len(r[:width])) for r in grid]
        merges = [(m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1) for m in ws.merged_cells.ranges]
        sheets.append(SheetData(ws.title, grid, formulas, formats, values, merges, ws.sheet_state != "visible"))
    return sheets


def _covered(merges) -> dict[tuple[int, int], tuple[int, int, int, int] | None]:
    """(row, col) -> merge range for origins, None for cells hidden under a merge."""
    out: dict = {}
    for r0, c0, r1, c1 in merges:
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                out[(r, c)] = None
        out[(r0, c0)] = (r0, c0, r1, c1)
    return out


@step("xlsx_html_semantic")
def xlsx_html(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    sheets = read_workbook(ctx, src.path)
    parts = []
    for sh in sheets:
        cover = _covered(sh.merges)
        rows = []
        for r in range(sh.rows):
            cells = []
            for c in range(sh.cols):
                m = cover.get((r, c), "free")
                if m is None:
                    continue
                attrs = ""
                if m != "free":
                    r0, c0, r1, c1 = m
                    r1, c1 = min(r1, sh.rows - 1), min(c1, sh.cols - 1)
                    attrs += f' rowspan="{r1 - r0 + 1}"' if r1 > r0 else ""
                    attrs += f' colspan="{c1 - c0 + 1}"' if c1 > c0 else ""
                key = (r, c)
                if key in sh.formulas:
                    attrs += f' data-formula="{html.escape(sh.formulas[key], quote=True)}"'
                v = sh.values.get(key)
                if isinstance(v, bool):
                    attrs += f' data-type="b" data-value="{str(v).upper()}"'
                elif isinstance(v, (int, float)):
                    attrs += f' data-type="n" data-value="{v!r}"'
                if key in sh.formats:
                    attrs += f' data-numfmt="{html.escape(sh.formats[key], quote=True)}"'
                text = html.escape(sh.shown[r][c]).replace("\n", "<br>")
                cells.append(f"<td{attrs}>{text}</td>")
            rows.append("<tr>" + "".join(cells) + "</tr>")
        hidden = ' class="hidden-sheet"' if sh.hidden else ""
        parts.append(f'<section{hidden}><h2>{html.escape(sh.name)}</h2>'
                     f'<table data-sheet="{html.escape(sh.name, quote=True)}">{"".join(rows)}</table></section>')
        ctx.added_text.append(sh.name)
    out = ctx.path("out.html")
    out.write_text(f'<!DOCTYPE html>\n<html><head><meta charset="utf-8"><title>{html.escape(ctx.source.stem)}</title>'
                   "<style>table{border-collapse:collapse}td{border:1px solid #bbb;padding:2px 6px;vertical-align:top}"
                   "</style></head><body>\n" + "\n".join(parts) + "\n</body></html>\n", encoding="utf-8")
    return Artifact("html", [out])


@step("xlsx_txt")
def xlsx_txt(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    sheets = read_workbook(ctx, src.path)
    chunks = []
    for sh in sheets:
        lines = [f"Sheet: {sh.name}"]
        ctx.added_text += ["Sheet:", sh.name]
        for row in sh.shown:
            cells = [v.replace("\t", " ").replace("\n", " ") for v in row]
            lines.append("\t".join(cells).rstrip("\t"))
        chunks.append("\n".join(lines))
    if any("\n" in v or "\t" in v for sh in sheets for row in sh.shown for v in row):
        ctx.expect("Line breaks and tabs inside cells became spaces (they separate rows and cells in the text).")
    out = ctx.path("out.txt")
    out.write_text("\n\n".join(chunks) + "\n", encoding="utf-8")
    return Artifact("txt", [out])


@step("xlsx_docx")
def xlsx_docx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt

    sheets = read_workbook(ctx, src.path)
    doc = Document()
    widest = max((sh.cols for sh in sheets), default=0)
    if widest > 6:
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for sh in sheets:
        ctx.check()
        doc.add_heading(sh.name, level=2)
        ctx.added_text.append(sh.name)
        if not sh.rows or not sh.cols:
            continue
        table = doc.add_table(rows=sh.rows, cols=sh.cols)
        table.style = "Table Grid"
        for r in range(sh.rows):
            for c in range(sh.cols):
                text = sh.shown[r][c]
                if text:
                    cell = table.cell(r, c)
                    cell.text = text
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(9)
        for r0, c0, r1, c1 in sh.merges:
            r1, c1 = min(r1, sh.rows - 1), min(c1, sh.cols - 1)
            if r1 > r0 or c1 > c0:
                keep = sh.shown[r0][c0]
                merged = table.cell(r0, c0).merge(table.cell(r1, c1))
                merged.text = keep
    out = ctx.path("out.docx")
    doc.save(out)
    ctx.expect("Formulas are shown as their results (Word tables hold values, not formulas).")
    return Artifact("docx", [out])


ROWS_PER_SLIDE, COLS_PER_SLIDE = 15, 8


@step("xlsx_pptx")
def xlsx_pptx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from pptx import Presentation
    from pptx.util import Emu, Pt

    sheets = read_workbook(ctx, src.path)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    layout = prs.slide_layouts[5]  # title only
    for sh in sheets:
        if not sh.rows or not sh.cols:
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = sh.name
            ctx.added_text.append(sh.name)
            continue
        for r_start in range(0, sh.rows, ROWS_PER_SLIDE):
            for c_start in range(0, sh.cols, COLS_PER_SLIDE):
                ctx.check()
                r_end = min(sh.rows, r_start + ROWS_PER_SLIDE)
                c_end = min(sh.cols, c_start + COLS_PER_SLIDE)
                slide = prs.slides.add_slide(layout)
                label = sh.name
                if sh.rows > ROWS_PER_SLIDE or sh.cols > COLS_PER_SLIDE:
                    label = f"{sh.name}: rows {r_start + 1}–{r_end}" + (
                        f", columns {_col(c_start)}–{_col(c_end - 1)}" if sh.cols > COLS_PER_SLIDE else "")
                slide.shapes.title.text = label
                ctx.added_text.append(label)
                nr, nc = r_end - r_start, c_end - c_start
                gf = slide.shapes.add_table(nr, nc, Emu(ooxml.emu(30)), Emu(ooxml.emu(110)), Emu(ooxml.emu(900)),
                                            Emu(ooxml.emu(min(400, 24 * nr))))
                tbl = gf.table
                for r in range(nr):
                    for c in range(nc):
                        cell = tbl.cell(r, c)
                        cell.text = sh.shown[r_start + r][c_start + c]
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(12)
                for r0, c0, r1, c1 in sh.merges:
                    a0, b0 = max(r0, r_start), max(c0, c_start)
                    a1, b1 = min(r1, r_end - 1), min(c1, c_end - 1)
                    if a0 <= a1 and b0 <= b1 and (a1 > a0 or b1 > b0):
                        keep = tbl.cell(a0 - r_start, b0 - c_start).text
                        tbl.cell(a0 - r_start, b0 - c_start).merge(tbl.cell(a1 - r_start, b1 - c_start))
                        tbl.cell(a0 - r_start, b0 - c_start).text = keep
    out = ctx.path("out.pptx")
    prs.save(out)
    ctx.expect("Formulas are shown as their results; large sheets are split across slides.")
    return Artifact("pptx", [out])


def _col(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s


@step("docx_xlsx")
def docx_xlsx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    body = ooxml_read.docx_body(src.path)
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    text_rows: list[tuple[str, str | None]] = []
    n = 0
    for kind, item in body:
        ctx.check()
        if kind == "table":
            n += 1
            g: ooxml_read.Grid = item  # type: ignore[assignment]
            ws = wb.create_sheet(safe_sheet_title(f"Table {n}", used))
            for (r, c), text in g.cells.items():
                cell = write_cell(ws, r + 1, c + 1, text)
                set_link(cell, g.links.get((r, c)))
                if "\n" in text:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            for r0, c0, r1, c1 in g.merges:
                ws.merge_cells(start_row=r0 + 1, start_column=c0 + 1, end_row=r1 + 1, end_column=c1 + 1)
        else:
            t = item.text  # type: ignore[union-attr]
            if t.strip():
                text_rows.append((t, next((r.link for r in item.runs if r.link and not r.link.startswith("#")), None)))
    if text_rows or not wb.sheetnames:
        ws = wb.create_sheet(safe_sheet_title("Text", used), 0)
        for t, link in text_rows:
            ws.append([None])
            c = write_cell(ws, ws.max_row, 1, t)
            set_link(c, link)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].width = 100
    out = ctx.path("out.xlsx")
    wb.save(out)
    ctx.expect("Each Word table is on its own sheet and the body text is on the 'Text' sheet, so the reading order changes.")
    return Artifact("xlsx", [out])

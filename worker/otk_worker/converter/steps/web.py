"""HTML, TXT, EPUB and SVG steps: Chromium print, Pandoc, and the toolkit's literal writers."""

from __future__ import annotations

import html
import math
import re
from pathlib import Path

from lxml import etree
from lxml import html as lhtml

from .. import chromium, engines, epub, htmlprep, ooxml, textutil
from ..context import Artifact, StepContext
from . import step
from .docmodel_out import safe_sheet_title, typed_value, write_cell


# ------------------------------------------------------------------ Chromium
@step("chromium_pdf")
def chromium_pdf(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    out = ctx.path("out.pdf")
    if src.format == "svg":
        w, h = htmlprep.svg_size_pt(src.path)
        svg = src.path.read_text(encoding="utf-8")
        svg = re.sub(r"^<\?xml[^>]*>\s*", "", svg.strip())
        svg = re.sub(r"<!DOCTYPE[^>]*>\s*", "", svg)
        page = ctx.path("page.html")
        page.write_text(
            f'<!DOCTYPE html><html><head><meta charset="utf-8"><style>@page{{size:{w:.3f}pt {h:.3f}pt;margin:0}}'
            f"html,body{{margin:0;padding:0}}body>svg{{display:block;width:{w:.3f}pt;height:{h:.3f}pt}}</style>"
            f"</head><body>{svg}</body></html>", encoding="utf-8")
        ctx.progress(0.1, "Chromium: printing the drawing")
        chromium.render_pdf(ctx, page, out, page_size_pt=(w, h))
    else:
        ctx.progress(0.1, "Chromium: printing the page")
        chromium.render_pdf(ctx, src.path, out)
    return Artifact("pdf", [out])


# ------------------------------------------------------------------ Pandoc
def _title_of(ctx: StepContext, src: Artifact) -> str:
    try:
        if src.format == "html":
            doc = lhtml.document_fromstring(src.path.read_text(encoding="utf-8"))
            t = doc.findtext(".//title")
            if t and t.strip():
                return t.strip()
            h1 = doc.find(".//h1")
            if h1 is not None and h1.text_content().strip():
                return h1.text_content().strip()
        elif src.format == "docx":
            from docx import Document
            t = Document(str(src.path)).core_properties.title
            if t:
                return t
        elif src.format == "epub":
            return epub.open_book(src.path).title or ctx.source.stem
    except Exception:
        pass
    return ctx.source.stem


def pandoc(ctx: StepContext, src: Path, fmt_in: str, out: Path, fmt_out: str, extra: list[str] | None = None) -> None:
    exe = engines.require("pandoc")
    cmd = [exe, "--sandbox", "--wrap=none", "-f", fmt_in, "-t", fmt_out, "-o", str(out)] + (extra or []) + [str(src)]
    res = engines.run(cmd, check=ctx.check, timeout=900, what="Pandoc", cwd=str(src.parent))
    for line in res.stderr.splitlines():
        if "[WARNING]" in line and "title" not in line.lower():
            ctx.note("Pandoc: " + line.replace("[WARNING]", "").strip())
    if not out.exists():
        raise engines.EngineFailed(f"Pandoc did not produce {out.suffix} output.")


_PANDOC_IN = {"html": "html", "docx": "docx", "epub": "epub"}


@step("html_docx_pandoc", "pandoc_to_docx")
def pandoc_docx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    out = ctx.path("out.docx")
    pandoc(ctx, src.path, _PANDOC_IN[src.format], out, "docx", ["-M", f"title={_title_of(ctx, src)}"])
    return Artifact("docx", [out])


@step("pandoc_to_pptx")
def pandoc_pptx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    out = ctx.path("out.pptx")
    pandoc(ctx, src.path, _PANDOC_IN[src.format], out, "pptx", ["-M", f"title={_title_of(ctx, src)}"])
    ctx.expect("Slides are cut at headings; page layout is not kept.")
    return Artifact("pptx", [out])


@step("pandoc_plain")
def pandoc_plain(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    out = ctx.path("out.txt")
    pandoc(ctx, src.path, _PANDOC_IN[src.format], out, "plain")
    return Artifact("txt", [out])


@step("pandoc_epub")
def pandoc_epub(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    out = ctx.path("out.epub")
    pandoc(ctx, src.path, _PANDOC_IN[src.format], out, "epub3", ["-M", f"title={_title_of(ctx, src)}", "--toc"])
    ctx.added_text.append(_title_of(ctx, src))
    return Artifact("epub", [out])


@step("pandoc_to_html")
def pandoc_html(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    out = ctx.path("out.html")
    pandoc(ctx, src.path, _PANDOC_IN[src.format], out, "html5",
           ["--standalone", "--embed-resources", "-M", f"title={_title_of(ctx, src)}"])
    return Artifact("html", [out])


# ------------------------------------------------------------------ HTML -> XLSX
def _html_grid(table) -> tuple[list[list], list[tuple[int, int, int, int]]]:
    """Cells of an HTML table as a grid (rowspan/colspan honoured) and the merged ranges (0-based)."""
    grid: dict[tuple[int, int], object] = {}
    merges = []
    r = 0
    for tr in table.iter("tr"):
        if tr.getparent() is not None and _closest_table(tr) is not table:
            continue
        c = 0
        for cell in tr:
            if not isinstance(cell.tag, str) or cell.tag.lower() not in ("td", "th"):
                continue
            while (r, c) in grid:
                c += 1
            rs = max(1, int(cell.get("rowspan") or 1)) if str(cell.get("rowspan") or "1").isdigit() else 1
            cs = max(1, int(cell.get("colspan") or 1)) if str(cell.get("colspan") or "1").isdigit() else 1
            for dr in range(rs):
                for dc in range(cs):
                    grid[(r + dr, c + dc)] = cell if (dr, dc) == (0, 0) else None
            if rs > 1 or cs > 1:
                merges.append((r, c, r + rs - 1, c + cs - 1))
            c += cs
        r += 1
    rows = max((k[0] for k in grid), default=-1) + 1
    cols = max((k[1] for k in grid), default=-1) + 1
    out = [[grid.get((i, j)) for j in range(cols)] for i in range(rows)]
    return out, merges


def _closest_table(el):
    p = el.getparent()
    while p is not None and p.tag != "table":
        p = p.getparent()
    return p


def _cell_text(cell) -> str:
    for br in cell.iter("br"):
        br.tail = "\n" + (br.tail or "")
    lines = textutil.html_lines(cell)
    return "\n".join(lines)


@step("html_xlsx")
def html_xlsx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    doc = lhtml.document_fromstring(src.path.read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    tables = [t for t in doc.iter("table") if _closest_table(t) is None]
    for ti, table in enumerate(tables, start=1):
        ctx.check()
        name = table.get("data-sheet") or (table.findtext("caption") or "").strip() or f"Table {ti}"
        ws = wb.create_sheet(safe_sheet_title(name, used))
        grid, merges = _html_grid(table)
        for r, row in enumerate(grid, start=1):
            for c, cell in enumerate(row, start=1):
                if cell is None:
                    continue
                formula = cell.get("data-formula")
                text = _cell_text(cell)
                if formula:
                    xc = ws.cell(row=r, column=c, value=formula if formula.startswith("=") else "=" + formula)
                else:
                    raw = cell.get("data-value")
                    if raw is not None and cell.get("data-type") == "n":
                        try:
                            xc = ws.cell(row=r, column=c, value=float(raw) if "." in raw or "e" in raw.lower() else int(raw))
                        except ValueError:
                            xc = write_cell(ws, r, c, text)
                    else:
                        xc = write_cell(ws, r, c, text)
                if cell.get("data-numfmt"):
                    xc.number_format = cell.get("data-numfmt")
                if "\n" in text:
                    xc.alignment = Alignment(wrap_text=True, vertical="top")
        for r0, c0, r1, c1 in merges:
            ws.merge_cells(start_row=r0 + 1, start_column=c0 + 1, end_row=r1 + 1, end_column=c1 + 1)
        cap = table.find("caption")
        if cap is not None:
            cap.getparent().remove(cap)
            ctx.note(f"The caption of table {ti} ('{cap.text_content().strip()[:60]}') is on the 'Text' sheet.")
            table.addprevious(cap)
            cap.tag = "p"
    body = doc.find("body")
    lines = textutil.html_lines(body if body is not None else doc, skip_tables=True)
    if lines or not wb.sheetnames:
        ws = wb.create_sheet(safe_sheet_title("Text", used))
        for line in lines:
            ws.append([None])
            write_cell(ws, ws.max_row, 1, line)
        ws.column_dimensions["A"].width = 100
    out = ctx.path("out.xlsx")
    wb.save(out)
    if tables:
        ctx.expect("Each table is on its own sheet and the other text is on the 'Text' sheet, so the reading order changes.")
    return Artifact("xlsx", [out])


# ------------------------------------------------------------------ TXT
def read_txt(ctx: StepContext, path: Path) -> str:
    text, enc = htmlprep.decode_bytes(path.read_bytes())
    if enc not in ("utf-8", "ascii"):
        ctx.note(f"The text file was read as {enc}.")
    return text.replace("\r\n", "\n").replace("\r", "\n")


@step("txt_html")
def txt_html(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    text = read_txt(ctx, src.path)
    lines = text.split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    body = []
    for ln in lines:
        if "\f" in ln:
            parts = ln.split("\f")
            for i, part in enumerate(parts):
                if i:
                    body.append('<hr class="form-feed">')
                if part or i == 0:
                    body.append(f'<p dir="auto">{html.escape(part) or "<br>"}</p>')
            continue
        body.append(f'<p dir="auto">{html.escape(ln) if ln else "<br>"}</p>')
    out = ctx.path("out.html")
    out.write_text(
        f'<!DOCTYPE html>\n<html><head><meta charset="utf-8"><title>{html.escape(ctx.source.stem)}</title><style>'
        "body{font-family:'Noto Sans Mono','DejaVu Sans Mono',Consolas,monospace;font-size:11pt}"
        "p{margin:0;white-space:pre-wrap;tab-size:8;min-height:1.2em}hr.form-feed{break-after:page;border:0}"
        "</style></head><body>\n" + "\n".join(body) + "\n</body></html>\n", encoding="utf-8")
    return Artifact("html", [out])


def _docx_line(doc, line: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = 0
    script = textutil.script_of(line)
    rtl = textutil.is_rtl_script(script)
    if rtl:
        p._p.get_or_add_pPr().insert(0, OxmlElement("w:bidi"))
    if not line:
        return
    run = p.add_run(line)  # tabs become <w:tab/>
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi"):
        fonts.set(qn(k), "Consolas")
    if script:
        fonts.set(qn("w:cs"), textutil.CS_FONT.get(script, "Arial"))
        fonts.set(qn("w:eastAsia"), textutil.CS_FONT.get(script, "Arial"))
        lang = OxmlElement("w:lang")
        if script in textutil.LANG_TAG:
            lang.set(qn("w:bidi"), textutil.LANG_TAG[script])
            rpr.append(lang)
    if rtl:
        rpr.append(OxmlElement("w:rtl"))


@step("txt_docx")
def txt_docx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    text = read_txt(ctx, src.path)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    lines = text.split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    for ln in lines:
        ctx.check()
        parts = ln.split("\f")
        for i, part in enumerate(parts):
            if i:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if part or i == 0:
                _docx_line(doc, part)
    doc.core_properties.title = ctx.source.stem
    out = ctx.path("out.docx")
    doc.save(out)
    return Artifact("docx", [out])


@step("txt_xlsx")
def txt_xlsx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    text = read_txt(ctx, src.path)
    wb = Workbook()
    ws = wb.active
    ws.title = safe_sheet_title(ctx.source.stem, set())
    lines = text.split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    bad = 0
    for r, ln in enumerate(lines, start=1):
        cells = ln.split("\t") if ctx.options.txt_split_tabs else [ln]
        for c, val in enumerate(cells, start=1):
            if not val:
                continue
            clean = ILLEGAL_CHARACTERS_RE.sub("�", val)
            bad += clean != val
            cell = ws.cell(row=r, column=c, value=clean)
            cell.number_format = "@"
            cell.data_type = "s"
    ws.column_dimensions["A"].width = 100 if not ctx.options.txt_split_tabs else 20
    if bad:
        ctx.expect(f"{bad} line(s) contained control characters Excel cannot store; they were replaced with �.")
    out = ctx.path("out.xlsx")
    wb.save(out)
    return Artifact("xlsx", [out])


def _fit_size(lines: list[str], width_pt: float, height_pt: float) -> float | None:
    for size in [x / 2 for x in range(48, 15, -1)]:  # 24pt down to 8pt
        rows = sum(max(1, math.ceil(len(ln.expandtabs(8)) * 0.6 * size / width_pt)) for ln in lines)
        if rows * size * 1.2 <= height_pt:
            return size
    return None


@step("txt_pptx")
def txt_pptx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu, Pt

    text = read_txt(ctx, src.path)
    if text.endswith("\n"):
        text = text[:-1]
    n = ctx.options.txt_lines_per_slide
    chunks: list[list[str]] = []
    for page in text.split("\f"):
        lines = page.split("\n")
        for i in range(0, max(1, len(lines)), n):
            chunks.append(lines[i:i + n])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    margin = 36.0
    box_w, box_h = 960 - 2 * margin, 540 - 2 * margin
    queue = list(chunks)
    slides = 0
    while queue:
        ctx.check()
        chunk = queue.pop(0)
        size = _fit_size(chunk, box_w, box_h)
        if size is None and len(chunk) > 1:
            half = len(chunk) // 2
            queue[0:0] = [chunk[:half], chunk[half:]]
            continue
        size = size or 8.0
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Emu(ooxml.emu(margin)), Emu(ooxml.emu(margin)), Emu(ooxml.emu(box_w)), Emu(ooxml.emu(box_h)))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        for i, ln in enumerate(chunk):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            script = textutil.script_of(ln)
            if textutil.is_rtl_script(script):
                para._p.get_or_add_pPr().set("rtl", "1")
            run = para.add_run()
            run.text = ln
            run.font.size = Pt(size)
            run.font.name = "Consolas"
            if script:
                from pptx.oxml import parse_xml
                rpr = run._r.get_or_add_rPr()
                rpr.append(parse_xml(f'<a:cs {ooxml.XMLNS} typeface="{textutil.CS_FONT.get(script, "Arial")}"/>'))
        slides += 1
    out = ctx.path("out.pptx")
    prs.save(out)
    if slides > 1:
        ctx.note(f"The text was split over {slides} slides (form feeds and every {n} lines); no line was cut.")
    return Artifact("pptx", [out])


# ------------------------------------------------------------------ EPUB -> HTML
@step("epub_html_join")
def epub_html_join(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    import posixpath
    from urllib.parse import unquote

    book = epub.open_book(src.path)
    zf = book.zf
    styles: list[str] = []
    seen_css: set[str] = set()
    sections: list[str] = []
    chapter_of = {p: i for i, p in enumerate(book.spine, start=1)}

    def css_inline(css: str, base: str) -> str:
        def url(m):
            ref = m.group(2).strip()
            if ref.startswith("data:") or htmlprep.REMOTE.match(ref):
                if htmlprep.REMOTE.match(ref):
                    ctx.note(f"Remote resource not loaded (offline): {ref}")
                    return 'url("")'
                return m.group(0)
            uri = epub.data_uri(zf, posixpath.normpath(posixpath.join(base, unquote(ref))))
            return f'url("{uri}")' if uri else 'url("")'
        return htmlprep._CSS_URL.sub(url, css)

    for i, zpath in enumerate(book.spine, start=1):
        ctx.check()
        base = posixpath.dirname(zpath)
        raw = zf.read(zpath)
        text, _ = htmlprep.decode_bytes(raw)
        root = lhtml.document_fromstring(text.encode("utf-8"), parser=lhtml.HTMLParser(encoding="utf-8"))
        for link in root.iter("link"):
            if "stylesheet" in (link.get("rel") or "") and link.get("href"):
                cpath = posixpath.normpath(posixpath.join(base, unquote(link.get("href"))))
                if cpath not in seen_css:
                    seen_css.add(cpath)
                    try:
                        styles.append(css_inline(zf.read(cpath).decode("utf-8", "replace"), posixpath.dirname(cpath)))
                    except KeyError:
                        ctx.note(f"Stylesheet {cpath} is missing from the EPUB.")
        for st in root.iter("style"):
            if st.text:
                styles.append(css_inline(st.text, base))
        body = root.find("body")
        if body is None:
            continue
        prefix = f"c{i}-"
        for el in body.iter():
            if not isinstance(el.tag, str):
                continue
            if el.get("id"):
                el.set("id", prefix + el.get("id"))
            for attr in ("src", "href", "{http://www.w3.org/1999/xlink}href", "poster"):
                ref = el.get(attr)
                if not ref:
                    continue
                tag = el.tag.split("}")[-1].lower()
                if tag == "a" and attr == "href":
                    if htmlprep.REMOTE.match(ref) or ref.startswith(("mailto:", "tel:")):
                        continue
                    doc_part, _, frag = ref.partition("#")
                    target_doc = posixpath.normpath(posixpath.join(base, unquote(doc_part))) if doc_part else zpath
                    if target_doc in chapter_of:
                        n = chapter_of[target_doc]
                        el.set("href", f"#c{n}-{frag}" if frag else f"#chapter-{n}")
                    continue
                if ref.startswith("#"):
                    el.set(attr, "#" + prefix + ref[1:])
                    continue
                if htmlprep.REMOTE.match(ref):
                    ctx.note(f"Remote resource not loaded (offline): {ref}")
                    el.set(attr, "")
                    continue
                uri = epub.data_uri(zf, posixpath.normpath(posixpath.join(base, unquote(ref))))
                if uri:
                    el.set(attr, uri)
            if el.get("style") and "url(" in el.get("style"):
                el.set("style", css_inline(el.get("style"), base))
        inner = "".join(lhtml.tostring(ch, encoding="unicode") for ch in body)
        if body.text and body.text.strip():
            inner = html.escape(body.text) + inner
        dir_attr = f' dir="{body.get("dir")}"' if body.get("dir") else ""
        sections.append(f'<section class="chapter" id="chapter-{i}"{dir_attr}>{inner}</section>')
    title = book.title or ctx.source.stem
    lang = f' lang="{html.escape(book.language)}"' if book.language else ""
    out = ctx.path("out.html")
    out.write_text(f"<!DOCTYPE html>\n<html{lang}><head><meta charset=\"utf-8\"><title>{html.escape(title)}</title>"
                   f"<style>{''.join(styles)}</style></head><body>\n" + "\n".join(sections) + "\n</body></html>\n",
                   encoding="utf-8")
    return Artifact("html", [out])


# ------------------------------------------------------------------ SVG
def _svg_markup(path: Path) -> str:
    s = path.read_text(encoding="utf-8")
    s = re.sub(r"^\s*<\?xml[^>]*>\s*", "", s)
    return re.sub(r"<!DOCTYPE[^>]*>\s*", "", s)


@step("svg_html")
def svg_html(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    w, h = htmlprep.svg_size_pt(src.path)
    title = ""
    root = etree.parse(str(src.path)).getroot()
    t = root.find("{http://www.w3.org/2000/svg}title")
    if t is not None and t.text:
        title = t.text.strip()
    out = ctx.path("out.html")
    out.write_text(f'<!DOCTYPE html>\n<html><head><meta charset="utf-8"><title>{html.escape(title or ctx.source.stem)}</title>'
                   f"<style>body{{margin:0}}body>svg{{display:block;width:{w:.2f}pt;height:{h:.2f}pt}}</style></head>"
                   f"<body>{_svg_markup(src.path)}</body></html>\n", encoding="utf-8")
    return Artifact("html", [out])


def svg_text_lines(path: Path) -> list[str]:
    root = etree.parse(str(path), etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True)).getroot()
    lines = []
    for el in root.iter():
        if not isinstance(el.tag, str):
            continue
        local = etree.QName(el).localname
        if local in ("title", "desc"):
            if el.text and el.text.strip():
                lines.append(el.text.strip())
        elif local == "text":
            t = re.sub(r"\s+", " ", "".join(el.itertext())).strip()
            if t:
                lines.append(t)
    return lines


@step("svg_txt")
def svg_txt(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    out = ctx.path("out.txt")
    out.write_text("\n".join(svg_text_lines(src.path)) + "\n", encoding="utf-8")
    return Artifact("txt", [out])


SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"


def _svg_blip(blip, rid: str) -> None:
    from pptx.oxml import parse_xml

    ext = parse_xml(f'<a:extLst {ooxml.XMLNS}><a:ext uri="{SVG_EXT_URI}"><asvg:svgBlip '
                    f'xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" r:embed="{rid}"/></a:ext></a:extLst>')
    blip.append(ext)


def _svg_png(ctx: StepContext, svg: Path) -> Path:
    from .raster import render_svg

    png = ctx.path("fallback.png")
    render_svg(ctx, svg, png, dpi=192)
    return png


@step("svg_pptx_picture")
def svg_pptx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml.ns import qn
    from pptx.util import Emu

    w, h = htmlprep.svg_size_pt(src.path)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(ooxml.emu(max(72, min(w, 4032)))), Emu(ooxml.emu(max(72, min(h, 4032))))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    png = _svg_png(ctx, src.path)
    s = min(prs.slide_width / ooxml.emu(w), prs.slide_height / ooxml.emu(h), 1.0)
    pic = slide.shapes.add_picture(str(png), 0, 0, Emu(int(ooxml.emu(w) * s)), Emu(int(ooxml.emu(h) * s)))
    part = Part(PackURI("/ppt/media/drawing1.svg"), "image/svg+xml", prs.part.package, src.path.read_bytes())
    rid = slide.part.relate_to(part, RT.IMAGE)
    _svg_blip(pic._element.find(".//" + qn("a:blip")), rid)
    out = ctx.path("out.pptx")
    prs.save(out)
    ctx.note("The drawing is a vector SVG picture with a PNG copy for older PowerPoint versions.")
    return Artifact("pptx", [out])


@step("svg_docx_picture")
def svg_docx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml.ns import qn
    from docx.shared import Pt

    w, h = htmlprep.svg_size_pt(src.path)
    doc = Document()
    sec = doc.sections[0]
    avail = (sec.page_width - sec.left_margin - sec.right_margin) / 12700
    s = min(1.0, avail / w)
    png = _svg_png(ctx, src.path)
    doc.add_picture(str(png), width=Pt(w * s))
    blip = doc.paragraphs[-1]._p.find(".//" + qn("a:blip"))
    part = Part(PackURI("/word/media/drawing1.svg"), "image/svg+xml", src.path.read_bytes(), doc.part.package)
    rid = doc.part.relate_to(part, RT.IMAGE)
    _svg_blip(blip, rid)
    out = ctx.path("out.docx")
    doc.save(out)
    if s < 1.0:
        ctx.note(f"The drawing was scaled to {s:.0%} to fit the page width.")
    ctx.note("The drawing is a vector SVG picture with a PNG copy for older Word versions.")
    return Artifact("docx", [out])

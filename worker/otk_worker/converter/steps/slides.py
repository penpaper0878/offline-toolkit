"""PowerPoint steps: slides to HTML, DOCX, XLSX and TXT (semantic, or rendered for exact output)."""

from __future__ import annotations

import base64
import html
import mimetypes
from io import BytesIO
from pathlib import Path

from .. import libreoffice as lo
from .. import ooxml_read, pdfpages
from ..context import Artifact, StepContext
from ..ooxml_read import DocxPara, SlideContent
from . import step
from .docmodel_out import safe_sheet_title, write_cell


def _runs_html(p: DocxPara) -> str:
    out = []
    for r in p.runs:
        if r.text == "\n":
            out.append("<br>")
            continue
        t = html.escape(r.text)
        if r.bold:
            t = f"<b>{t}</b>"
        if r.italic:
            t = f"<i>{t}</i>"
        if r.link:
            t = f'<a href="{html.escape(r.link, quote=True)}">{t}</a>'
        out.append(t)
    return "".join(out)


def _grid_html(g: ooxml_read.Grid) -> str:
    covered = set()
    spans = {}
    for r0, c0, r1, c1 in g.merges:
        spans[(r0, c0)] = (r1 - r0 + 1, c1 - c0 + 1)
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                if (r, c) != (r0, c0):
                    covered.add((r, c))
    rows = []
    for r in range(g.rows):
        cells = []
        for c in range(g.cols):
            if (r, c) in covered:
                continue
            rs, cs = spans.get((r, c), (1, 1))
            attrs = (f' rowspan="{rs}"' if rs > 1 else "") + (f' colspan="{cs}"' if cs > 1 else "")
            cells.append(f"<td{attrs}>{html.escape(g.cells.get((r, c), '')).replace(chr(10), '<br>')}</td>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return "<table>" + "".join(rows) + "</table>"


def _data_uri(blob: bytes, ext: str) -> str:
    mime = mimetypes.guess_type(f"x.{ext}")[0] or "image/png"
    return f"data:{mime};base64,{base64.b64encode(blob).decode()}"


def _slide_label(s: SlideContent, ctx: StepContext) -> str:
    if s.title:
        return s.title
    label = f"Slide {s.index}"
    ctx.added_text.append(label)
    return label


def _paras_html(paras: list[DocxPara]) -> str:
    out = []
    in_list = False
    for p in paras:
        if not p.text.strip():
            continue
        d = ' dir="rtl"' if p.rtl else ""
        if p.list_level:
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f'<li{d} style="margin-left:{p.list_level}em">{_runs_html(p)}</li>')
        else:
            if in_list:
                out.append("</ul>")
                in_list = False
            out.append(f"<p{d}>{_runs_html(p)}</p>")
    if in_list:
        out.append("</ul>")
    return "".join(out)


@step("pptx_html_semantic")
def pptx_html_semantic(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    slides = ooxml_read.pptx_slides(src.path)
    parts = []
    for s in slides:
        body = [f"<h2>{html.escape(_slide_label(s, ctx))}</h2>"]
        for it in s.items:
            if it.is_title:
                continue
            if it.kind == "text":
                body.append(_paras_html(it.paragraphs))
            elif it.kind == "table":
                body.append(_grid_html(it.grid))
            elif it.kind == "picture":
                w, h = it.size_pt
                body.append(f'<figure><img src="{_data_uri(it.image, it.image_ext)}" alt="" '
                            f'style="width:{w:.0f}pt;max-width:100%;height:auto"></figure>')
        if s.notes:
            body.append('<aside class="notes"><h3>Speaker notes</h3>'
                        + "".join(f"<p>{html.escape(n)}</p>" for n in s.notes if n.strip()) + "</aside>")
            ctx.added_text += ["Speaker", "notes"]
        parts.append(f'<section class="slide" id="slide-{s.index}">' + "".join(body) + "</section>")
    out = ctx.path("out.html")
    out.write_text(f'<!DOCTYPE html>\n<html><head><meta charset="utf-8"><title>{html.escape(ctx.source.stem)}</title>'
                   "<style>body{font-family:'Noto Sans',Arial,sans-serif;max-width:60em;margin:2em auto;padding:0 1em}"
                   ".slide{border-bottom:1px solid #ccc;padding-bottom:1em}table{border-collapse:collapse}"
                   "td{border:1px solid #888;padding:3px 6px}aside.notes{background:#f5f5f5;padding:.5em 1em}"
                   "</style></head><body>\n" + "\n".join(parts) + "\n</body></html>\n", encoding="utf-8")
    return Artifact("html", [out])


def _render_slides(ctx: StepContext, src: Path) -> Path:
    ctx.progress(0.05, "LibreOffice: rendering the slides")
    return lo.convert(src, ctx.folder("lo"), lo.pdf_filter("impress"), check=ctx.check)


@step("pptx_html_exact")
def pptx_html_exact(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    import pymupdf

    from .pdf import extract_model

    pdf = _render_slides(ctx, src.path)
    model, _ = extract_model(ctx, pdf)
    slides = ooxml_read.pptx_slides(src.path)
    parts = []
    with pymupdf.open(pdf) as doc:
        for i, page in enumerate(doc):
            svg = pdfpages.uniquify_ids(pdfpages.page_svg(page, model, exact=True), f"s{i + 1}-")
            notes = ""
            if i < len(slides) and slides[i].notes:
                notes = ('<aside class="notes">' + "".join(f"<p>{html.escape(n)}</p>" for n in slides[i].notes if n.strip())
                         + "</aside>")
            parts.append(f'<section class="slide" id="slide-{i + 1}">'
                         f"{pdfpages.inline_svg(svg, page.rect.width, page.rect.height)}{notes}</section>")
    out = ctx.path("out.html")
    out.write_text(f'<!DOCTYPE html>\n<html><head><meta charset="utf-8"><title>{html.escape(ctx.source.stem)}</title>'
                   "<style>body{background:#e8e8e8;margin:0;padding:16px;font-family:'Noto Sans',Arial,sans-serif}"
                   ".slide{margin:0 auto 24px;width:max-content}.slide svg{display:block;background:#fff;"
                   "box-shadow:0 1px 4px rgba(0,0,0,.3)}aside.notes{background:#fff;padding:.5em 1em;margin-top:6px}"
                   "</style></head><body>\n" + "\n".join(parts) + "\n</body></html>\n", encoding="utf-8")
    ctx.note("Slides are drawn as they appear in LibreOffice; the speaker notes are under each slide.")
    return Artifact("html", [out])


def _docx_paragraph(doc, p: DocxPara, style: str | None = None):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    par = doc.add_paragraph(style=style)
    if p.rtl:
        par._p.get_or_add_pPr().insert(0, OxmlElement("w:bidi"))
    for r in p.runs:
        if r.text == "\n":
            par.add_run().add_break()
            continue
        if r.link:
            rid = doc.part.relate_to(r.link, RT.HYPERLINK, is_external=True)
            h = OxmlElement("w:hyperlink")
            h.set(qn("r:id"), rid)
            run = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            st = OxmlElement("w:rStyle")
            st.set(qn("w:val"), "Hyperlink")
            rpr.append(st)
            run.append(rpr)
            t = OxmlElement("w:t")
            t.text = r.text
            t.set(qn("xml:space"), "preserve")
            run.append(t)
            h.append(run)
            par._p.append(h)
        else:
            run = par.add_run(r.text)
            run.bold = r.bold or None
            run.italic = r.italic or None
    return par


def _docx_grid(doc, g: ooxml_read.Grid) -> None:
    if not g.rows or not g.cols:
        return
    table = doc.add_table(rows=g.rows, cols=g.cols)
    table.style = "Table Grid"
    for (r, c), text in g.cells.items():
        table.cell(r, c).text = text
    for r0, c0, r1, c1 in g.merges:
        keep = g.cells.get((r0, c0), "")
        table.cell(r0, c0).merge(table.cell(r1, c1)).text = keep


@step("pptx_docx_semantic")
def pptx_docx_semantic(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from docx import Document
    from docx.shared import Pt

    slides = ooxml_read.pptx_slides(src.path)
    doc = Document()
    sec = doc.sections[0]
    avail = (sec.page_width - sec.left_margin - sec.right_margin) / 12700
    for s in slides:
        ctx.check()
        h = doc.add_heading(_slide_label(s, ctx), level=1)
        if s.index > 1:
            h.paragraph_format.page_break_before = True
        for it in s.items:
            if it.is_title:
                continue
            if it.kind == "text":
                for p in it.paragraphs:
                    if p.text.strip():
                        _docx_paragraph(doc, p, "List Bullet" if p.list_level else None)
            elif it.kind == "table":
                _docx_grid(doc, it.grid)
            elif it.kind == "picture":
                width = min(avail, it.size_pt[0] or avail)
                try:
                    doc.add_picture(BytesIO(it.image), width=Pt(max(1.0, width)))
                except Exception:
                    ctx.note(f"A picture on slide {s.index} ({it.image_ext}) could not be placed in Word.")
        if s.notes and any(n.strip() for n in s.notes):
            doc.add_heading("Speaker notes", level=3)
            ctx.added_text += ["Speaker", "notes"]
            for n in s.notes:
                if n.strip():
                    doc.add_paragraph(n).runs[0].italic = True
    out = ctx.path("out.docx")
    doc.save(out)
    ctx.expect("Slide layout is not kept: each slide becomes a heading followed by its text, tables and pictures.")
    return Artifact("docx", [out])


@step("pptx_docx_handout")
def pptx_docx_handout(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    import pymupdf
    from docx import Document
    from docx.shared import Pt

    pdf = _render_slides(ctx, src.path)
    slides = ooxml_read.pptx_slides(src.path)
    doc = Document()
    sec = doc.sections[0]
    avail = (sec.page_width - sec.left_margin - sec.right_margin) / 12700
    img_dir = ctx.folder("slides")
    with pymupdf.open(pdf) as d:
        for i, page in enumerate(d):
            ctx.progress(0.2 + 0.7 * i / max(1, d.page_count), f"Slide {i + 1}")
            png = img_dir / f"{i + 1:04d}.png"
            page.get_pixmap(dpi=220, alpha=False).save(png)
            pic_par = doc.add_paragraph()
            if i:
                pic_par.paragraph_format.page_break_before = True
            run = pic_par.add_run()
            inline = run.add_picture(str(png), width=Pt(avail))
            content = slides[i] if i < len(slides) else None
            if content is not None:
                alt = " ".join(p.text for it in content.items if it.kind == "text" for p in it.paragraphs)
                inline._inline.docPr.set("descr", alt[:2000])
                for n in content.notes:
                    if n.strip():
                        doc.add_paragraph(n)
    out = ctx.path("out.docx")
    doc.save(out)
    ctx.expect_check("text", "Slide text is part of the slide pictures in a handout (it is in each picture's alt text); "
                             "only the speaker notes are document text.")
    return Artifact("docx", [out])


@step("pptx_xlsx")
def pptx_xlsx(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    slides = ooxml_read.pptx_slides(src.path)
    wb = Workbook()
    ws = wb.active
    used = {"slides"}
    ws.title = "Slides"
    ws.append(["Slide", "Text", "Notes"])
    ctx.added_text += ["Slide", "Text", "Notes"]
    for s in slides:
        texts = []
        for it in s.items:
            if it.kind == "text":
                texts.extend(p.text for p in it.paragraphs if p.text.strip())
        ws.append([s.index, None, None])
        ctx.added_text.append(str(s.index))
        write_cell(ws, ws.max_row, 2, "\n".join(texts)).alignment = Alignment(wrap_text=True, vertical="top")
        write_cell(ws, ws.max_row, 3, "\n".join(n for n in s.notes if n.strip())).alignment = Alignment(wrap_text=True, vertical="top")
        for k, it in enumerate([i for i in s.items if i.kind == "table"], start=1):
            tws = wb.create_sheet(safe_sheet_title(f"Slide {s.index} table {k}", used))
            for (r, c), text in it.grid.cells.items():
                write_cell(tws, r + 1, c + 1, text)
            for r0, c0, r1, c1 in it.grid.merges:
                tws.merge_cells(start_row=r0 + 1, start_column=c0 + 1, end_row=r1 + 1, end_column=c1 + 1)
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 50
    out = ctx.path("out.xlsx")
    wb.save(out)
    ctx.expect("Slide tables are on their own sheets; slide text and notes are listed on the 'Slides' sheet.")
    return Artifact("xlsx", [out])


@step("pptx_txt")
def pptx_txt(ctx: StepContext, src: Artifact, target: str) -> Artifact:
    slides = ooxml_read.pptx_slides(src.path)
    chunks = []
    for s in slides:
        lines = [f"Slide {s.index}"]
        ctx.added_text += ["Slide", str(s.index)]
        for it in s.items:
            if it.kind == "text":
                lines.extend(p.text for p in it.paragraphs if p.text.strip())
            elif it.kind == "table":
                g = it.grid
                for r in range(g.rows):
                    lines.append("\t".join(g.cells.get((r, c), "").replace("\n", " ") for c in range(g.cols)).rstrip("\t"))
        notes = [n for n in s.notes if n.strip()]
        if notes:
            lines.append("Notes:")
            ctx.added_text.append("Notes:")
            lines.extend(notes)
        chunks.append("\n".join(lines))
    out = ctx.path("out.txt")
    out.write_text("\n\n".join(chunks) + "\n", encoding="utf-8")
    return Artifact("txt", [out])

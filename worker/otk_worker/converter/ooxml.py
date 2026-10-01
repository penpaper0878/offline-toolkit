"""DrawingML building blocks shared by the DOCX and PPTX writers.

Positions come in PDF points (top-left origin) and are converted to EMU.
Vector paths become <a:custGeom> (lines and cubic Béziers kept exactly),
text becomes native text boxes, pictures keep their original bytes.
"""

from __future__ import annotations

import itertools
from xml.sax.saxutils import escape

EMU_PER_PT = 12700

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
}
XMLNS = " ".join(f'xmlns:{k}="{v}"' for k, v in NS.items())

_ids = itertools.count(1000)


def next_id() -> int:
    return next(_ids)


def emu(pt: float) -> int:
    return int(round(pt * EMU_PER_PT))


def attr(s: str) -> str:
    return escape(s, {'"': "&quot;"})


def color_xml(hex_color: str, alpha: float = 1.0) -> str:
    val = hex_color.lstrip("#").upper()
    a = "" if alpha >= 0.999 else f'<a:alpha val="{int(round(max(0.0, alpha) * 100000))}"/>'
    return f'<a:srgbClr val="{val}">{a}</a:srgbClr>'


def custgeom(shape, box: tuple[float, float, float, float]) -> str:
    """<a:custGeom> for a docmodel Shape, in coordinates local to `box` (x0, y0, x1, y1 in pt)."""
    x0, y0, x1, y1 = box
    w, h = max(1, emu(x1 - x0)), max(1, emu(y1 - y0))

    def pt(x, y):
        return f'<a:pt x="{max(0, min(w, emu(x - x0)))}" y="{max(0, min(h, emu(y - y0)))}"/>'

    parts = []
    for op in shape.ops:
        k = op[0]
        if k == "m":
            parts.append(f"<a:moveTo>{pt(op[1], op[2])}</a:moveTo>")
        elif k == "l":
            parts.append(f"<a:lnTo>{pt(op[1], op[2])}</a:lnTo>")
        elif k == "c":
            parts.append(f"<a:cubicBezTo>{pt(op[1], op[2])}{pt(op[3], op[4])}{pt(op[5], op[6])}</a:cubicBezTo>")
        elif k == "h":
            parts.append("<a:close/>")
    fill_attr = "" if shape.fill else ' fill="none"'
    stroke_attr = "" if shape.stroke else ' stroke="0"'
    return (f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
            f'<a:pathLst><a:path w="{w}" h="{h}"{fill_attr}{stroke_attr}>{"".join(parts)}</a:path></a:pathLst>'
            f'</a:custGeom>')


def shape_props(shape, box) -> str:
    """xfrm + geometry + fill + line for a docmodel Shape (used inside p:spPr / wps:spPr)."""
    x0, y0, x1, y1 = box
    fill = (f"<a:solidFill>{color_xml(shape.fill, shape.fill_opacity)}</a:solidFill>" if shape.fill else "<a:noFill/>")
    if shape.stroke:
        dash = ""
        if shape.dashes:
            dash = "<a:custDash>" + "".join(
                f'<a:ds d="{int(shape.dashes[i] / max(shape.width, 0.1) * 100000)}" '
                f'sp="{int(shape.dashes[(i + 1) % len(shape.dashes)] / max(shape.width, 0.1) * 100000)}"/>'
                for i in range(0, len(shape.dashes), 2)) + "</a:custDash>"
        cap = {0: "flat", 1: "rnd", 2: "sq"}.get(shape.cap, "flat")
        join = {0: '<a:miter lim="800000"/>', 1: "<a:round/>", 2: "<a:bevel/>"}.get(shape.join, "")
        line = (f'<a:ln w="{emu(max(shape.width, 0.25))}" cap="{cap}"><a:solidFill>{color_xml(shape.stroke, shape.stroke_opacity)}'
                f"</a:solidFill>{dash}{join}</a:ln>")
    else:
        line = "<a:ln><a:noFill/></a:ln>"
    return (f'<a:xfrm><a:off x="{emu(x0)}" y="{emu(y0)}"/><a:ext cx="{max(1, emu(x1 - x0))}" cy="{max(1, emu(y1 - y0))}"/></a:xfrm>'
            f"{custgeom(shape, box)}{fill}{line}")


def shape_box(shape) -> tuple[float, float, float, float]:
    xs = [v for op in shape.ops for v in op[1::2]]
    ys = [v for op in shape.ops for v in op[2::2]]
    if not xs:
        return tuple(shape.bbox)  # type: ignore[return-value]
    return min(xs), min(ys), max(xs), max(ys)


# ------------------------------------------------------------------ PPTX
def pptx_shape_xml(shape) -> str:
    box = shape_box(shape)
    sid = next_id()
    return (f'<p:sp {XMLNS}><p:nvSpPr><p:cNvPr id="{sid}" name="Shape {sid}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f"<p:spPr>{shape_props(shape, box)}</p:spPr></p:sp>")


def pptx_add_shape(slide, shape) -> None:
    from pptx.oxml import parse_xml

    slide.shapes._spTree.append(parse_xml(pptx_shape_xml(shape)))


def pptx_run_props(run, span, invisible: bool = False) -> None:
    """Font, size, style, colour and script fonts for a python-pptx run."""
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    from . import fontnames

    fam, bold, italic = fontnames.split(span.font)
    f = run.font
    f.size = Pt(max(1.0, round(span.size * 2) / 2))
    f.bold = span.bold or bold
    f.italic = span.italic or italic
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        for old in rpr.findall(qn(tag)):
            rpr.remove(old)
    from lxml import etree

    if invisible or span.alpha < 0.01:
        for old in rpr.findall(qn("a:solidFill")):
            rpr.remove(old)
        # Built in place (not parsed from a string) so no namespace declarations are repeated on every run.
        fill = etree.Element(qn("a:solidFill"))
        clr = etree.SubElement(fill, qn("a:srgbClr"), val="000000")
        etree.SubElement(clr, qn("a:alpha"), val="0")
        rpr.insert(0, fill)
    else:
        f.color.rgb = RGBColor.from_string(span.color.lstrip("#").upper())
    for tag in ("latin", "ea", "cs"):
        etree.SubElement(rpr, qn(f"a:{tag}"), typeface=fam)
    if span.link:
        run.hyperlink.address = span.link


def pptx_text_box(slide, lines, box, *, wrap: bool, rtl: bool, angle: float = 0.0, invisible: bool = False,
                  align: str | None = None):
    """A text box with one paragraph per entry in `lines` (each a list of spans)."""
    from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu

    x0, y0, x1, y1 = box
    tb = slide.shapes.add_textbox(Emu(emu(x0)), Emu(emu(y0)), Emu(max(1, emu(x1 - x0))), Emu(max(1, emu(y1 - y0))))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if angle:
        tb.rotation = -angle
    for i, spans in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_before = para.space_after = 0
        if rtl:
            para._p.get_or_add_pPr().set("rtl", "1")
            para.alignment = PP_ALIGN.RIGHT
        elif align == "center":
            para.alignment = PP_ALIGN.CENTER
        for span in spans:
            if not span.text:
                continue
            run = para.add_run()
            run.text = span.text
            pptx_run_props(run, span, invisible)
    bodypr = tf._txBody.find(qn("a:bodyPr"))
    if bodypr is not None and not wrap:
        bodypr.set("wrap", "none")
    return tb


# ------------------------------------------------------------------ DOCX
def _anchor(inner: str, box, *, z: int, behind: bool, name: str, graphic_uri: str, rot: float = 0.0) -> str:
    x0, y0, x1, y1 = box
    did = next_id()
    return (f'<wp:anchor {XMLNS} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{max(2, z)}" '
            f'behindDoc="{1 if behind else 0}" locked="0" layoutInCell="1" allowOverlap="1">'
            f'<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>{emu(x0)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{emu(y0)}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{max(1, emu(x1 - x0))}" cy="{max(1, emu(y1 - y0))}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
            f'<wp:wrapNone/><wp:docPr id="{did}" name="{attr(name)} {did}"/><wp:cNvGraphicFramePr/>'
            f'<a:graphic><a:graphicData uri="{graphic_uri}">{inner}</a:graphicData></a:graphic></wp:anchor>')


WPS_URI = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
PIC_URI = "http://schemas.openxmlformats.org/drawingml/2006/picture"


def _wrap_run(drawing: str, alternate: bool) -> str:
    body = f"<w:drawing>{drawing}</w:drawing>"
    if alternate:
        body = f'<mc:AlternateContent><mc:Choice Requires="wps">{body}</mc:Choice><mc:Fallback/></mc:AlternateContent>'
    return f'<w:r {XMLNS}><w:rPr><w:noProof/></w:rPr>{body}</w:r>'


def docx_picture_run(rid: str, box, *, z: int, behind: bool = True, name: str = "Picture") -> str:
    x0, y0, x1, y1 = box
    pid = next_id()
    inner = (f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{pid}" name="{attr(name)}"/><pic:cNvPicPr/></pic:nvPicPr>'
             f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
             f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{max(1, emu(x1 - x0))}" cy="{max(1, emu(y1 - y0))}"/></a:xfrm>'
             f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>')
    return _wrap_run(_anchor(inner, box, z=z, behind=behind, name=name, graphic_uri=PIC_URI), alternate=False)


def docx_shape_run(shape, *, z: int) -> str:
    box = shape_box(shape)
    props = shape_props(shape, box).replace(f'<a:off x="{emu(box[0])}" y="{emu(box[1])}"/>', '<a:off x="0" y="0"/>')
    inner = f'<wps:wsp><wps:cNvSpPr/><wps:spPr>{props}</wps:spPr><wps:bodyPr/></wps:wsp>'
    return _wrap_run(_anchor(inner, box, z=z, behind=True, name="Shape", graphic_uri=WPS_URI), alternate=True)


def docx_run_xml(span, *, rtl: bool, invisible: bool = False, rid: str | None = None) -> str:
    from . import fontnames

    fam, bold, italic = fontnames.split(span.font)
    size = max(2, int(round(span.size * 2)))
    props = [f'<w:rFonts w:ascii="{attr(fam)}" w:hAnsi="{attr(fam)}" w:eastAsia="{attr(fam)}" w:cs="{attr(fam)}"/>']
    if span.bold or bold:
        props += ["<w:b/>", "<w:bCs/>"]
    if span.italic or italic:
        props += ["<w:i/>", "<w:iCs/>"]
    if invisible or span.alpha < 0.01:
        props.append('<w:color w:val="FFFFFF"/>')
    else:
        props.append(f'<w:color w:val="{span.color.lstrip("#").upper()}"/>')
    props += [f'<w:sz w:val="{size}"/>', f'<w:szCs w:val="{size}"/>']
    if rtl:
        props.append("<w:rtl/>")
    if invisible or span.alpha < 0.01:
        props.append('<w14:textFill><w14:noFill/></w14:textFill>')
    text = escape(span.text)
    run = f'<w:r><w:rPr>{"".join(props)}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'
    if rid:
        run = f'<w:hyperlink r:id="{rid}">{run}</w:hyperlink>'
    return run


def docx_textbox_run(paragraphs_xml: list[str], box, *, z: int, behind: bool = False, rot: float = 0.0,
                     wrap: bool = False) -> str:
    x0, y0, x1, y1 = box
    rot_attr = f' rot="{int(round(-rot * 60000)) % 21600000}"' if rot else ""
    inner = (f'<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm{rot_attr}><a:off x="0" y="0"/>'
             f'<a:ext cx="{max(1, emu(x1 - x0))}" cy="{max(1, emu(y1 - y0))}"/></a:xfrm>'
             f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>'
             f'<wps:txbx><w:txbxContent>{"".join(paragraphs_xml)}</w:txbxContent></wps:txbx>'
             f'<wps:bodyPr rot="0" vert="horz" wrap="{"square" if wrap else "none"}" lIns="0" tIns="0" rIns="0" bIns="0" '
             f'anchor="t" anchorCtr="0"><a:noAutofit/></wps:bodyPr></wps:wsp>')
    return _wrap_run(_anchor(inner, box, z=z, behind=behind, name="Text", graphic_uri=WPS_URI), alternate=True)


def docx_para_xml(runs_xml: str, *, rtl: bool, align: str | None = None) -> str:
    ppr = ['<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>']
    if rtl:
        ppr.insert(0, "<w:bidi/>")
    if align:  # right-to-left paragraphs start at the right edge by default
        ppr.append(f'<w:jc w:val="{align}"/>')
    return f'<w:p><w:pPr>{"".join(ppr)}</w:pPr>{runs_xml}</w:p>'

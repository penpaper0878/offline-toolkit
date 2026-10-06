"""Independent extractors: what a file contains, read with a different library than the one that wrote it.

PDF: pypdfium2 (text, outline, rendering) + pikepdf (links, images, fonts) —
the conversions read PDFs with PyMuPDF, so a PyMuPDF bug cannot hide itself.
DOCX/PPTX/XLSX: direct XML / python-docx / python-pptx / openpyxl; XLS: xlrd.
HTML/EPUB/SVG: lxml. Images: Pillow.
"""

from __future__ import annotations

import base64
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import numpy as np
from lxml import etree
from lxml import html as lhtml
from PIL import Image as PILImage

from .. import fontnames, ooxml_read, textutil

PILImage.MAX_IMAGE_PIXELS = 300_000_000


@dataclass
class Img:
    width: int
    height: int
    digest: str                      # SHA-1 of the decoded RGB(A) pixels
    thumb: np.ndarray | None = None  # 32x32 grey, for "visually the same" matching


@dataclass
class Table:
    rows: int
    cols: int
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)


@dataclass
class Extract:
    format: str
    text: str = ""
    repeatable: set[str] = field(default_factory=set)        # tokens that may repeat (headers, footers, page numbers)
    has_lists: bool = False
    units: dict[str, int] = field(default_factory=dict)       # pages / slides / sheets
    images: list[Img] | None = None                          # None = not checked for this format
    tables: list[Table] | None = None
    links: set[str] | None = None
    bookmarks: list[str] | None = None
    notes: list[str] | None = None
    fonts_requested: set[str] | None = None                  # family names the document asks for
    fonts_used: dict[str, bool] | None = None                # family -> embedded (PDF) / referenced (Office)
    cells: dict[str, dict[str, tuple]] | None = None         # sheet -> {"A1": (shown, formula, numfmt)}
    merges: dict[str, list] | None = None                    # sheet -> merged ranges
    render: Callable[[int], list[PILImage.Image]] | None = None   # page renders at the given DPI
    second_text: Callable[[], str] | None = None             # a second PDF reader, used when the first disagrees
    info: list[str] = field(default_factory=list)


def _img(im: PILImage.Image) -> Img:
    im.load()
    if im.mode not in ("RGB", "RGBA", "L", "LA"):
        im = im.convert("RGBA" if "A" in im.getbands() or im.mode == "P" and "transparency" in im.info else "RGB")
    if im.mode == "LA":
        im = im.convert("RGBA")
    if im.mode == "L":
        im = im.convert("RGB")
    digest = hashlib.sha1(im.tobytes()).hexdigest()
    thumb = np.asarray(im.convert("L").resize((32, 32), PILImage.Resampling.BOX), dtype=np.float32)
    return Img(im.width, im.height, digest, thumb)


def _img_bytes(data: bytes) -> Img | None:
    try:
        return _img(PILImage.open(io.BytesIO(data)))
    except Exception:
        return None


# ------------------------------------------------------------------ PDF
def _pdfium_render(path: Path, password: str | None = None):
    def render(dpi: int) -> list[PILImage.Image]:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(path), password=password)
        out = []
        try:
            for i in range(len(doc)):
                out.append(doc[i].render(scale=dpi / 72, draw_annots=True).to_pil().convert("RGB"))
        finally:
            doc.close()
        return out
    return render


def pdf(path: Path, fmt: str = "pdf", password: str | None = None) -> Extract:
    import pikepdf
    import pypdfium2 as pdfium
    from pikepdf import PdfImage

    ex = Extract(fmt)
    doc = pdfium.PdfDocument(str(path), password=password)
    texts = []
    try:
        ex.units["pages"] = len(doc)
        for i in range(len(doc)):
            page = doc[i]
            tp = page.get_textpage()
            texts.append(tp.get_text_range())
            tp.close()
            page.close()
        ex.bookmarks = [b.get_title() for b in doc.get_toc()]
    finally:
        doc.close()
    ex.text = "\n".join(t.replace("\r\n", "\n").replace("\r", "\n") for t in texts)
    ex.links, ex.images, ex.fonts_used = set(), [], {}
    with pikepdf.open(path, password=password or "") as pdfdoc:
        for page in pdfdoc.pages:
            for annot in page.obj.get("/Annots", []) or []:
                try:
                    if annot.get("/Subtype") == "/Link":
                        act = annot.get("/A")
                        if act is not None and act.get("/S") == "/URI":
                            ex.links.add(str(act.get("/URI")))
                except Exception:
                    continue
            for raw in _drawn_images(page):
                try:
                    pim = PdfImage(raw)
                    pil = pim.as_pil_image()
                    if raw.get("/SMask") is not None:
                        mask = PdfImage(raw.SMask).as_pil_image().convert("L")
                        if mask.size == pil.size:
                            pil = pil.convert("RGB")
                            pil.putalpha(mask)
                    got = _img(pil)
                    ex.images.append(got)
                except Exception:
                    ex.info.append("An image in the PDF could not be decoded for comparison.")
            res = page.obj.get("/Resources", {})
            fonts = res.get("/Font", {}) if res is not None else {}
            for _, font in (fonts.items() if hasattr(fonts, "items") else []):
                try:
                    _pdf_font(font, ex.fonts_used)
                except Exception:
                    continue
    ex.fonts_requested = {f for f in ex.fonts_used if not f.startswith("Type3")}
    ex.render = _pdfium_render(path, password)

    def mupdf_text() -> str:
        from ..docmodel import text_only
        return text_only(path, password)
    ex.second_text = mupdf_text
    return ex


def _drawn_images(page) -> list:
    """Image XObjects the page really paints (shared resource dictionaries list images of other pages too)."""
    import pikepdf

    out: list = []
    seen_forms: set = set()

    def walk(stream_owner, resources, depth=0):
        if depth > 8 or resources is None:
            return
        xobjs = resources.get("/XObject", {}) if hasattr(resources, "get") else {}
        try:
            ops = pikepdf.parse_content_stream(stream_owner)
        except Exception:
            return
        for operands, op in ops:
            if str(op) != "Do" or not operands:
                continue
            xo = xobjs.get(str(operands[0])) if hasattr(xobjs, "get") else None
            if xo is None:
                continue
            sub = xo.get("/Subtype")
            if sub == "/Image":
                out.append(xo)
            elif sub == "/Form":
                key = xo.objgen
                if key in seen_forms and key != (0, 0):
                    continue
                seen_forms.add(key)
                walk(xo, xo.get("/Resources", resources), depth + 1)

    walk(page, page.obj.get("/Resources"))
    return out


def _pdf_font(font, out: dict[str, bool]) -> None:
    base = str(font.get("/BaseFont", "")).lstrip("/")
    desc = font.get("/FontDescriptor")
    if font.get("/Subtype") == "/Type0":
        kids = font.get("/DescendantFonts")
        if kids:
            desc = kids[0].get("/FontDescriptor")
    if font.get("/Subtype") == "/Type3":
        out[f"Type3 ({base or 'unnamed'})"] = True
        return
    embedded = bool(desc is not None and any(k in desc for k in ("/FontFile", "/FontFile2", "/FontFile3")))
    fam = fontnames.family(base)
    out[fam] = out.get(fam, False) or embedded


# ------------------------------------------------------------------ DOCX
def _docx_theme_fonts(z: zipfile.ZipFile) -> dict[str, str]:
    try:
        theme = etree.fromstring(z.read("word/theme/theme1.xml"))
    except KeyError:
        return {}
    A = ooxml_read.A
    out = {}
    for kind, key in (("majorFont", "major"), ("minorFont", "minor")):
        el = theme.find(f".//{{{A}}}{kind}")
        if el is not None:
            latin = el.find(f"{{{A}}}latin")
            if latin is not None:
                out[f"{key}HAnsi"] = out[f"{key}Ascii"] = latin.get("typeface")
            cs = el.find(f"{{{A}}}cs")
            if cs is not None and cs.get("typeface"):
                out[f"{key}Bidi"] = cs.get("typeface")
    return out


_CS_TEXT = re.compile("[\u0590-\u08ff\u0900-\u0dff\ufb1d-\ufdff\ufe70-\ufefc]")
_EA_TEXT = re.compile("[\u3040-\u30ff\u3400-\u9fff\uac00-\ud7af\uf900-\ufaff]")
_LATIN_TEXT = re.compile(r"[A-Za-z0-9\u00c0-\u024f]")
_FONT_KEYS = ("ascii", "hAnsi", "cs", "eastAsia")


def _rfonts(rpr, W: str) -> dict[str, str]:
    """{'ascii': name or 'theme:minorHAnsi', ...} from an rPr element (theme attributes win)."""
    out: dict[str, str] = {}
    if rpr is None:
        return out
    rf = rpr.find(f"{{{W}}}rFonts")
    if rf is None:
        return out
    for k in _FONT_KEYS:
        theme = rf.get(f"{{{W}}}{k}Theme")
        name = rf.get(f"{{{W}}}{k}")
        if theme:
            out[k] = "theme:" + theme
        elif name:
            out[k] = name
    return out


def _docx_fonts(z: zipfile.ZipFile) -> set[str]:
    """Fonts the text actually uses: per run, the Latin font for Latin text, the complex-script font for
    Indic/Arabic/Hebrew text, the East Asian font for CJK; resolved through run, character and paragraph
    styles (with basedOn) down to the document defaults."""
    W = ooxml_read.W
    theme = _docx_theme_fonts(z)
    styles: dict[str, tuple[str | None, dict]] = {}
    defaults: dict[str, str] = {}
    default_para = None
    try:
        sroot = etree.fromstring(z.read("word/styles.xml"))
        dd = sroot.find(f"{{{W}}}docDefaults/{{{W}}}rPrDefault/{{{W}}}rPr")
        defaults = _rfonts(dd, W)
        for st in sroot.iter(f"{{{W}}}style"):
            sid = st.get(f"{{{W}}}styleId")
            based = st.find(f"{{{W}}}basedOn")
            styles[sid] = (based.get(f"{{{W}}}val") if based is not None else None, _rfonts(st.find(f"{{{W}}}rPr"), W))
            if st.get(f"{{{W}}}type") == "paragraph" and st.get(f"{{{W}}}default") in ("1", "true"):
                default_para = sid
    except KeyError:
        pass

    def chain(sid: str | None) -> list[dict]:
        out, seen = [], set()
        while sid and sid in styles and sid not in seen:
            seen.add(sid)
            out.append(styles[sid][1])
            sid = styles[sid][0]
        return out

    def resolve(value: str | None) -> str | None:
        if value and value.startswith("theme:"):
            t = value[6:]
            return theme.get(t) or theme.get(t.replace("Ascii", "HAnsi"))
        return value

    names: set[str] = set()
    root = etree.fromstring(z.read("word/document.xml"))
    for p in root.iter(f"{{{W}}}p"):
        if ooxml_read._in_fallback(p):
            continue
        ppr = p.find(f"{{{W}}}pPr")
        pstyle = None
        if ppr is not None and ppr.find(f"{{{W}}}pStyle") is not None:
            pstyle = ppr.find(f"{{{W}}}pStyle").get(f"{{{W}}}val")
        para_layers = chain(pstyle or default_para)
        for r in p.findall(f".//{{{W}}}r"):
            if ooxml_read._in_fallback(r) or r.getparent() is None:
                continue
            text = "".join(t.text or "" for t in r.findall(f"{{{W}}}t"))
            if not text.strip():
                continue
            rpr = r.find(f"{{{W}}}rPr")
            layers = [_rfonts(rpr, W)]
            if rpr is not None and rpr.find(f"{{{W}}}rStyle") is not None:
                layers += chain(rpr.find(f"{{{W}}}rStyle").get(f"{{{W}}}val"))
            layers += para_layers + [defaults]

            def pick(key: str) -> str | None:
                for layer in layers:
                    if key in layer:
                        return resolve(layer[key])
                return None
            if _LATIN_TEXT.search(text):
                f = pick("ascii") or pick("hAnsi")
                if f:
                    names.add(f)
            if _CS_TEXT.search(text):
                f = pick("cs")
                if f:
                    names.add(f)
            if _EA_TEXT.search(text):
                f = pick("eastAsia")
                if f:
                    names.add(f)
    return {n for n in names if n}


def docx(path: Path) -> Extract:
    ex = Extract("docx")
    body = ooxml_read.docx_body(path)
    parts, ex.tables, ex.links, ex.bookmarks = [], [], set(), []
    for kind, item in body:
        if kind == "p":
            parts.append(item.text)
            if item.level:
                ex.bookmarks.append(item.text.strip())
            if item.list_level is not None:
                ex.has_lists = True
            for r in item.runs:
                if r.link and not r.link.startswith("#"):
                    ex.links.add(r.link)
        else:
            g: ooxml_read.Grid = item
            ex.tables.append(Table(g.rows, g.cols, sorted(g.merges)))
            for r in range(g.rows):
                parts.append(" ".join(g.cells.get((r, c), "") for c in range(g.cols)))
    W = ooxml_read.W
    ex.images = []
    with zipfile.ZipFile(path) as z:
        rels = {}
        try:
            for rel in etree.fromstring(z.read("word/_rels/document.xml.rels")):
                if rel.get("TargetMode") != "External":
                    rels[rel.get("Id")] = "word/" + rel.get("Target").lstrip("/").replace("word/", "", 1) \
                        if not rel.get("Target").startswith("/") else rel.get("Target").lstrip("/")
        except KeyError:
            pass
        doc = etree.fromstring(z.read("word/document.xml"))
        for blip in doc.iter(f"{{{ooxml_read.A}}}blip"):
            if ooxml_read._in_fallback(blip):
                continue
            rid = blip.get(f"{{{ooxml_read.R}}}embed")
            target = rels.get(rid)
            if target:
                try:
                    got = _img_bytes(z.read(target))
                except KeyError:
                    got = None
                if got:
                    ex.images.append(got)
        # Footnotes and endnotes are document text; headers and footers repeat on every page.
        for name in sorted(z.namelist()):
            if re.match(r"word/(footnotes|endnotes)\.xml$", name):
                root = etree.fromstring(z.read(name))
                for note in root:
                    if note.get(f"{{{W}}}type") in ("separator", "continuationSeparator", "continuationNotice"):
                        continue
                    t = "".join(x.text or "" for x in note.iter(f"{{{W}}}t"))
                    if t.strip():
                        parts.append(t)
            elif re.match(r"word/(header|footer)\d*\.xml$", name):
                root = etree.fromstring(z.read(name))
                t = " ".join(x.text or "" for x in root.iter(f"{{{W}}}t"))
                ex.repeatable.update(t.split())
                if any((ins.text or "").strip().upper().startswith(("PAGE", "NUMPAGES")) for ins in root.iter(f"{{{W}}}instrText")) \
                        or any("PAGE" in (f.get(f"{{{W}}}instr") or "").upper() for f in root.iter(f"{{{W}}}fldSimple")):
                    ex.repeatable.update(str(i) for i in range(1, 1000))
        ex.fonts_requested = _docx_fonts(z)
    ex.fonts_used = {f: True for f in ex.fonts_requested}
    ex.text = "\n".join(parts)
    return ex


# ------------------------------------------------------------------ PPTX
def pptx(path: Path) -> Extract:
    ex = Extract("pptx")
    slides = ooxml_read.pptx_slides(path)
    ex.units["slides"] = len(slides)
    parts, ex.tables, ex.images, ex.links, ex.notes = [], [], [], set(), []
    ex.bookmarks = None  # PowerPoint has no bookmarks (slide titles are text, checked with the text)
    for s in slides:
        for it in s.items:
            if it.kind == "text":
                for p in it.paragraphs:
                    parts.append(p.text)
                    if p.list_level:
                        ex.has_lists = True
                    for r in p.runs:
                        if r.link and not r.link.startswith("#"):
                            ex.links.add(r.link)
            elif it.kind == "table":
                g = it.grid
                ex.tables.append(Table(g.rows, g.cols, sorted(g.merges)))
                for r in range(g.rows):
                    parts.append(" ".join(g.cells.get((r, c), "") for c in range(g.cols)))
            elif it.kind == "picture" and it.image:
                got = _img_bytes(it.image)
                if got:
                    ex.images.append(got)
        ex.notes.extend(n for n in s.notes if n.strip())
    ex.text = "\n".join(parts)
    ex.fonts_requested = _pptx_fonts(path)
    ex.fonts_used = {f: True for f in ex.fonts_requested}
    return ex


def _pptx_fonts(path: Path) -> set[str]:
    A = ooxml_read.A
    names: set[str] = set()
    theme: dict[str, str] = {}
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("ppt/theme/theme") and n.endswith(".xml"):
                root = etree.fromstring(z.read(n))
                for kind, key in (("majorFont", "+mj-lt"), ("minorFont", "+mn-lt")):
                    el = root.find(f".//{{{A}}}{kind}/{{{A}}}latin")
                    if el is not None and key not in theme:
                        theme[key] = el.get("typeface")
        for n in z.namelist():
            if re.match(r"ppt/slides/slide\d+\.xml$", n):
                root = etree.fromstring(z.read(n))
                used_text = False
                for r in root.iter(f"{{{A}}}r"):
                    used_text = True
                    latin = r.find(f"{{{A}}}rPr/{{{A}}}latin")
                    if latin is not None and latin.get("typeface"):
                        tf = latin.get("typeface")
                        names.add(theme.get(tf, tf))
                if used_text:
                    names.update(v for v in theme.values() if v)
    return {n for n in names if n and not n.startswith("+")}


# ------------------------------------------------------------------ XLSX / XLS
def xlsx(path: Path, shown: dict[str, list[list[str]]] | None) -> Extract:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    ex = Extract("xlsx")
    wb = load_workbook(path, data_only=False)
    ex.units["sheets"] = len(wb.worksheets)
    ex.cells, ex.merges, ex.tables, ex.links, ex.images = {}, {}, [], set(), []
    parts = []
    fonts = set()
    for ws in wb.worksheets:
        grid = (shown or {}).get(ws.title)
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                coord = f"{get_column_letter(c.column)}{c.row}"
                formula = str(c.value) if c.data_type == "f" else None
                text = None
                if grid is not None and c.row - 1 < len(grid) and c.column - 1 < len(grid[c.row - 1]):
                    text = grid[c.row - 1][c.column - 1]
                cells[coord] = (text if text is not None else ("" if formula else str(c.value)), formula,
                                c.number_format if c.number_format != "General" else None)
                if c.hyperlink is not None and c.hyperlink.target:
                    ex.links.add(c.hyperlink.target)
                if c.font is not None and c.font.name:
                    fonts.add(c.font.name)
        ex.cells[ws.title] = cells
        ex.merges[ws.title] = sorted(str(m) for m in ws.merged_cells.ranges)
        ex.tables.append(Table(ws.max_row if cells else 0, ws.max_column if cells else 0,
                               sorted((m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1)
                                      for m in ws.merged_cells.ranges)))
        for img in getattr(ws, "_images", []):
            try:
                got = _img_bytes(img._data())
                if got:
                    ex.images.append(got)
            except Exception:
                pass
        if grid is not None:
            parts.extend(" ".join(v for v in row if v) for row in grid)
        else:
            parts.extend(v[0] for v in cells.values() if v[0])
    ex.text = "\n".join(parts)
    ex.fonts_requested = fonts
    ex.fonts_used = {f: True for f in fonts}
    return ex


def xls(path: Path, shown: dict[str, list[list[str]]] | None) -> Extract:
    import xlrd

    ex = Extract("xls")
    book = xlrd.open_workbook(str(path), formatting_info=True)
    ex.units["sheets"] = book.nsheets
    ex.cells, ex.merges, ex.tables = {}, {}, []
    parts = []
    for sh in book.sheets():
        grid = (shown or {}).get(sh.name)
        cells = {}
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                coord = f"{_col(c)}{r + 1}"
                xf = book.xf_list[cell.xf_index] if cell.xf_index is not None else None
                fmt = book.format_map[xf.format_key].format_str if xf is not None and xf.format_key in book.format_map else None
                text = grid[r][c] if grid is not None and r < len(grid) and c < len(grid[r]) else str(cell.value)
                cells[coord] = (text, None, fmt if fmt and fmt != "General" else None)
        ex.cells[sh.name] = cells
        merges = sorted((r0, c0, r1 - 1, c1 - 1) for r0, r1, c0, c1 in sh.merged_cells)
        ex.merges[sh.name] = [f"{_col(c0)}{r0 + 1}:{_col(c1)}{r1 + 1}" for r0, c0, r1, c1 in merges]
        ex.tables.append(Table(sh.nrows, sh.ncols, merges))
        if grid is not None:
            parts.extend(" ".join(v for v in row if v) for row in grid)
    ex.text = "\n".join(parts)
    ex.info.append("XLS formulas are compared by their results (xlrd does not read formula text).")
    return ex


def _col(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s


# ------------------------------------------------------------------ HTML / EPUB / SVG / TXT
def _html_tables(root) -> list[Table]:
    from ..steps.web import _closest_table, _html_grid

    out = []
    for t in root.iter("table"):
        if _closest_table(t) is not None:
            continue
        grid, merges = _html_grid(t)
        out.append(Table(len(grid), max((len(r) for r in grid), default=0), sorted(merges)))
    return out


def _data_image(src: str) -> Img | None:
    m = re.match(r"data:[^;,]+(;base64)?,(.*)$", src or "", re.S)
    if not m or not m.group(1):
        return None
    try:
        return _img_bytes(base64.b64decode(m.group(2)))
    except Exception:
        return None


def html_doc(root, fmt: str = "html") -> Extract:
    ex = Extract(fmt)
    body = root.find(".//body")
    if body is None:
        body = root
    ex.text = "\n".join(textutil.html_lines(body))
    ex.tables = _html_tables(body)
    ex.links = {a.get("href") or a.get("xlink:href") for a in body.iter("a")
                if (a.get("href") or a.get("xlink:href") or "").startswith(("http:", "https:", "mailto:"))}
    ex.bookmarks = [re.sub(r"\s+", " ", h.text_content()).strip() for h in body.iter("h1", "h2", "h3") if h.text_content().strip()]
    for nav in body.iter("nav"):
        ex.bookmarks += [re.sub(r"\s+", " ", a.text_content()).strip() for a in nav.iter("a") if a.text_content().strip()]
    ex.has_lists = any(True for _ in body.iter("ol", "ul"))
    ex.images = []
    for img in body.iter("img", "image"):
        got = _data_image(img.get("src") or img.get("href") or img.get("xlink:href") or "")
        ex.images.append(got if got else Img(0, 0, "unreadable"))
    fams = set()
    for style in root.iter("style"):
        for m in re.finditer(r"font-family\s*:\s*([^;}]+)", style.text or ""):
            fams.add(m.group(1))
    for el in root.iter():
        st = el.get("style") if isinstance(el.tag, str) else None
        if st and "font-family" in st:
            for m in re.finditer(r"font-family\s*:\s*([^;]+)", st):
                fams.add(m.group(1))
    ex.fonts_requested = {f for f in fams}
    return ex


def html(path: Path) -> Extract:
    root = lhtml.document_fromstring(path.read_text(encoding="utf-8", errors="replace"))
    return html_doc(root)


def epub(path: Path) -> Extract:
    from .. import epub as epubmod

    book = epubmod.open_book(path)
    ex = Extract("epub")
    texts, ex.tables, ex.links, ex.images = [], [], set(), []
    fixed = False
    try:
        opf_names = [n for n in book.zf.namelist() if n.endswith(".opf")]
        fixed = any(b"pre-paginated" in book.zf.read(n) for n in opf_names)
    except Exception:
        pass
    for zpath in book.spine:
        if zpath == book.nav_path:
            continue  # the table of contents page: its entries are bookmarks, not body text
        try:
            root = lhtml.document_fromstring(book.zf.read(zpath))
        except Exception:
            continue
        sub = html_doc(root, "epub")
        texts.append(sub.text)
        ex.tables.extend(sub.tables or [])
        ex.links.update(sub.links or set())
        base = zpath.rsplit("/", 1)[0] if "/" in zpath else ""
        for img in root.iter("img", "{http://www.w3.org/2000/svg}image", "image"):
            src = (img.get("src") or img.get("{http://www.w3.org/1999/xlink}href") or img.get("xlink:href")
                   or img.get("href") or "")
            got = _data_image(src)
            if got is None and src and not src.startswith(("http:", "https:")):
                import posixpath
                from urllib.parse import unquote
                try:
                    got = _img_bytes(book.zf.read(posixpath.normpath(posixpath.join(base, unquote(src)))))
                except KeyError:
                    got = None
            if got is not None:
                ex.images.append(got)
        if fixed:
            # Fixed layout: the visible page is the picture; the text layer is in invisible SVG <text>.
            svg_text = [re.sub(r"\s+", " ", "".join(t.itertext())).strip()
                        for t in root.iter("{http://www.w3.org/2000/svg}text", "text")]
            if svg_text and not sub.text.strip():
                texts[-1] = "\n".join(x for x in svg_text if x)
    ex.text = "\n".join(texts)
    ex.bookmarks = [t for t, _ in book.nav if t]
    if fixed:
        ex.units["pages"] = len(book.spine)
    return ex


def svg(paths: list[Path]) -> Extract:
    from ..steps.web import svg_text_lines

    ex = Extract("svg")
    parts, ex.links, ex.images = [], set(), []
    for p in paths:
        parts.extend(svg_text_lines(p, visible_only=True))
        root = etree.parse(str(p), etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True)).getroot()
        for el in root.iter():
            if not isinstance(el.tag, str):
                continue
            local = etree.QName(el).localname
            href = el.get("{http://www.w3.org/1999/xlink}href") or el.get("href") or ""
            if local == "a" and href.startswith(("http:", "https:", "mailto:")):
                ex.links.add(href)
            if local == "image":
                got = _data_image(href)
                if got:
                    ex.images.append(got)
    ex.text = "\n".join(parts)
    if len(paths) > 1:
        ex.units["pages"] = len(paths)
    ex.render = _svg_render(paths)
    return ex


def _svg_render(paths: list[Path]):
    """MuPDF draws SVG (an independent renderer from resvg and Chromium); scaled to the SVG's physical size."""
    def render(dpi: int) -> list[PILImage.Image]:
        import pymupdf

        from ..htmlprep import svg_size_pt
        out = []
        for p in paths:
            w_pt, _ = svg_size_pt(p)
            with pymupdf.open(p) as d:
                page = d[0]
                zoom = (w_pt * dpi / 72) / max(1.0, page.rect.width)
                pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
                out.append(PILImage.frombytes("RGB", (pix.width, pix.height), pix.samples))
        return out
    return render


def txt(path: Path) -> Extract:
    from ..htmlprep import decode_bytes

    text, _ = decode_bytes(path.read_bytes())
    return Extract("txt", text=text.replace("\r\n", "\n"))


def _on_white(im: PILImage.Image) -> PILImage.Image:
    """RGB as the picture is seen: transparent areas over white (the page every other renderer uses), not black."""
    if "A" in im.getbands() or (im.mode == "P" and "transparency" in im.info):
        rgba = im.convert("RGBA")
        flat = PILImage.new("RGB", rgba.size, (255, 255, 255))
        flat.paste(rgba, mask=rgba.getchannel("A"))
        return flat
    return im.convert("RGB")


def raster(paths: list[Path], fmt: str) -> Extract:
    ex = Extract(fmt, images=[])
    dpis = []
    for p in paths:
        im = PILImage.open(p)
        dpis.append(im.info.get("dpi"))
        ex.images.append(_img(im))
    ex.units["pages"] = len(paths)

    def render(dpi: int) -> list[PILImage.Image]:
        out = []
        for p, d in zip(paths, dpis):
            im = _on_white(PILImage.open(p))
            src_dpi = float(d[0]) if d else 300.0
            s = dpi / src_dpi
            out.append(im.resize((max(1, round(im.width * s)), max(1, round(im.height * s))), PILImage.Resampling.BOX))
        return out
    ex.render = render
    return ex

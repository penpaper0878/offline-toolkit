"""Post-conversion verification: extract both sides independently, compare, decide the verdict."""

from __future__ import annotations

import csv
from pathlib import Path

from .. import libreoffice as lo
from ..context import Artifact, StepContext
from ..planner import Catalog, Route
from . import compare, extract

PDF_LIKE = ("pdf", "pdfa1b", "pdfa2b", "pdfa3b", "pdf_scanned")


def _shown(ctx: StepContext, path: Path, names: list[str]) -> dict[str, list[list[str]]]:
    files = lo.sheet_csvs(path, ctx.folder("verify-csv"), names, check=ctx.check)
    out = {}
    for name, f in files.items():
        with open(f, encoding="utf-8", newline="") as fh:
            out[name] = [list(r) for r in csv.reader(fh, delimiter="\t", quotechar='"')]
    return out


def _sheet_names(path: Path, fmt: str) -> list[str]:
    if fmt == "xls":
        import xlrd
        return xlrd.open_workbook(str(path), on_demand=True).sheet_names()
    from openpyxl import load_workbook
    return load_workbook(path, read_only=True).sheetnames


def extract_file(ctx: StepContext, paths: list[Path], fmt: str) -> extract.Extract:
    p = paths[0]
    if fmt in PDF_LIKE:
        return extract.pdf(p, fmt)
    if fmt == "docx":
        return extract.docx(p)
    if fmt == "pptx":
        return extract.pptx(p)
    if fmt in ("xlsx", "xls"):
        try:
            shown = _shown(ctx, p, _sheet_names(p, fmt))
        except Exception as exc:  # LibreOffice missing: compare raw values only
            ctx.note(f"Displayed cell values could not be read ({exc}); raw values were compared.")
            shown = None
        return extract.xlsx(p, shown) if fmt == "xlsx" else extract.xls(p, shown)
    if fmt in ("doc", "ppt"):
        modern = lo.convert(p, ctx.folder("verify-lo"), lo.DOCX if fmt == "doc" else lo.PPTX, check=ctx.check)
        ex = extract.docx(modern) if fmt == "doc" else extract.pptx(modern)
        ex.format = fmt
        ex.info.append(f"The {fmt.upper()} source was read through LibreOffice (no independent reader for this format).")
        return ex
    if fmt == "html":
        return extract.html(p)
    if fmt == "epub":
        return extract.epub(p)
    if fmt == "txt":
        return extract.txt(p)
    if fmt == "svg":
        return extract.svg(paths)
    if fmt in ("png", "jpeg"):
        return extract.raster(paths, fmt)
    raise ValueError(f"no extractor for {fmt}")


def _chromium_svg_render(ctx: StepContext, svg: Path):
    """Reference rendering of an SVG source: Chromium prints it, pdfium draws the page."""
    def render(dpi: int):
        from ..context import Artifact as Art
        from ..steps.web import chromium_pdf
        pdf = chromium_pdf(ctx, Art("svg", [svg]), "pdf")
        return extract._pdfium_render(pdf.path)(dpi)
    return render


def _office_render(ctx: StepContext, path: Path, fmt: str):
    """Render a DOCX/PPTX output through LibreOffice for the appearance check."""
    def render(dpi: int):
        family = "writer" if fmt == "docx" else "impress"
        pdf = lo.convert(path, ctx.folder("verify-render"), lo.pdf_filter(family), check=ctx.check)
        return extract._pdfium_render(pdf)(dpi)
    return render


def verify_conversion(ctx: StepContext, *, source: Path, source_format: str, output: Artifact, target: str,
                      route: Route, validation, catalog: Catalog) -> dict:
    src_path = source
    if source_format == "html" and (ctx.work / "source-prepared.html").exists():
        src_path = ctx.work / "source-prepared.html"
    elif source_format == "svg" and (ctx.work / "source-prepared.svg").exists():
        src_path = ctx.work / "source-prepared.svg"
    notes: list[str] = []
    try:
        src = extract_file(ctx, [src_path], source_format)
    except Exception as exc:
        src = None
        notes.append(f"The source could not be read for verification: {exc}")
    if src is not None and source_format == "svg":
        if route.steps and route.steps[0].engine == "chromium" or output.info.get("engine") == "chromium":
            ctx.expected_checks["appearance_skip"] = ("The reference renderer for SVG is Chromium, which also made "
                                                      "this output, so a comparison would prove nothing.")
        else:
            src.render = _chromium_svg_render(ctx, src_path)
    out_fmt = target if target not in ("pdfa1b", "pdfa2b", "pdfa3b") else "pdf"
    try:
        out = extract_file(ctx, output.paths, out_fmt)
    except Exception as exc:
        out = None
        notes.append(f"The output could not be read for verification: {exc}")
    lost = {loss.feature: loss.level for loss in catalog.potential_losses(route)}
    sit = compare.Situation(
        source_format=source_format, target=target, mode=route.mode, lost=lost,
        target_caps=catalog.formats[target]["caps"], expected_checks=dict(ctx.expected_checks),
        added_text=list(ctx.added_text), expected_notes=list(ctx.expected),
        ocr_route=any(s.engine in ("ocr", "ocrmypdf") for s in route.steps)
        or (ctx.options.ocr and (source_format in ("png", "jpeg", "pdf_scanned") or bool(ctx.scanned_pages))),
        notes_pages=ctx.options.notes_pages)
    checks: list[compare.Check] = []
    if src is not None and out is not None:
        for fn in (compare.check_text, compare.check_units, compare.check_images, compare.check_tables,
                   compare.check_links, compare.check_bookmarks, compare.check_notes, compare.check_cells,
                   compare.check_fonts):
            try:
                checks.append(fn(src, out, sit))
            except Exception as exc:
                checks.append(compare.Check(fn.__name__.replace("check_", ""), fn.__name__.replace("check_", "").title(),
                                            compare.FAIL, f"The check could not run: {exc}"))
        if ctx.options.verify_appearance and route.mode == "exact":
            out_render = None
            if out_fmt in ("docx", "pptx") and src.render is not None:
                out_render = _office_render(ctx, output.path, out_fmt)
            try:
                checks.append(compare.check_appearance(src, out, sit, out_render))
            except Exception as exc:
                checks.append(compare.Check("appearance", "Appearance (rendered pages)", compare.FAIL,
                                            f"The pages could not be rendered for comparison: {exc}"))
        notes += src.info + out.info
    else:
        checks.append(compare.Check("read", "Verification", compare.FAIL, "Verification could not run; see the notes."))
    if validation is not None:
        if validation.compliant:
            checks.append(compare.Check("pdfa", f"PDF/A-{validation.flavour} (veraPDF)", compare.PASS,
                                        "veraPDF: the file complies.", validation.to_dict()))
        else:
            checks.append(compare.Check("pdfa", f"PDF/A-{validation.flavour} (veraPDF)", compare.FAIL,
                                        validation.message if not validation.available else
                                        f"veraPDF: {len(validation.failed_rules)} rule(s) failed.", validation.to_dict()))
    if ctx.low_confidence:
        checks.append(compare.Check("ocr", "OCR confidence", compare.EXPECTED,
                                    f"{len(ctx.low_confidence)} word(s) below {ctx.options.ocr_min_confidence:.0%} confidence.",
                                    {"words": ctx.low_confidence[:300]}))
    v = compare.verdict(checks, ctx.expected)
    counts: dict[str, int] = {}
    for c in checks:
        counts[c.status] = counts.get(c.status, 0) + 1
    from .. import formats as fmts
    return {
        "verdict": v,
        "source": {"name": ctx.source.name, "path": str(ctx.source), "format": source_format,
                   "label": catalog.formats[source_format]["label"]},
        "target": {"format": target, "label": catalog.formats[target]["label"], "ext": fmts.OUTPUT_EXT[target]},
        "mode": route.mode,
        "route": {"chain": route.chain, "steps": [{"edge": s.id, "engine": catalog.engines.get(s.engine, {}).get("label", s.engine),
                                                   "summary": s.summary} for s in route.steps],
                  "override": route.override_reason},
        "plannedLosses": [{"feature": loss.feature, "label": loss.label, "level": loss.level}
                          for loss in catalog.potential_losses(route)],
        "stepLosses": catalog.step_losses(route),
        "checks": [c.to_dict() for c in checks],
        "summary": counts,
        "expectedChanges": list(ctx.expected),
        "notes": list(dict.fromkeys(ctx.notes + notes)),
        "options": ctx.options.to_dict(),
    }

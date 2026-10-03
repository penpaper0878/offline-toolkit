"""PowerPoint: every layer a native, separately editable object on one slide.

Background picture -> the slide background (stays put while editing); photos -> pictures; shapes ->
preset shapes (rectangle, rounded rectangle, ellipse, line); graphics -> a group of freeform shapes, one
per colour, so each can be recoloured; text -> text boxes with the matched font, size, colour, style and
alignment; tables -> native tables with merged cells, fills and rules.

Text placement: with exact line spacing (what this writer sets) a presentation text frame puts the first
baseline one line pitch minus 0.2 em below its top, for every font (measured in LibreOffice, which
follows PowerPoint's font-independent line layout here). The frame is placed from that rule.
"""

from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from ..converter.ooxml import XMLNS, attr, next_id
from . import fonts_out, layout
from .vectorize import parse_path

EMU_PER_PT = 12700
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
PPTX_FIRST_BASELINE = 0.2      # em above the end of the first line pitch


class _Units:
    def __init__(self, k: float):
        self.k = k   # points per page pixel

    def emu(self, px: float) -> int:
        return int(round(px * self.k * EMU_PER_PT))

    def pt(self, px: float) -> float:
        return px * self.k


def _clr(hex_color: str, alpha: float = 1.0) -> str:
    a = "" if alpha >= 0.999 else f'<a:alpha val="{int(round(max(0.0, alpha) * 100000))}"/>'
    return f'<a:srgbClr val="{hex_color.lstrip("#").upper()}">{a}</a:srgbClr>'


def _xfrm(u: _Units, x, y, w, h, rot: float = 0.0, flip_h=False, flip_v=False, tag="a:xfrm") -> str:
    r = f' rot="{int(round(rot * 60000)) % 21600000}"' if rot else ""
    f = (' flipH="1"' if flip_h else "") + (' flipV="1"' if flip_v else "")
    return (f'<{tag}{r}{f}><a:off x="{u.emu(x)}" y="{u.emu(y)}"/><a:ext cx="{max(1, u.emu(w))}" cy="{max(1, u.emu(h))}"/>'
            f'</{tag}>')


def _line_xml(u: _Units, lyr: dict) -> str:
    if lyr.get("stroke") and lyr.get("strokeWidth"):
        return (f'<a:ln w="{max(1, int(round(u.pt(lyr["strokeWidth"]) * EMU_PER_PT)))}">'
                f'<a:solidFill>{_clr(lyr["stroke"], lyr.get("opacity", 1.0))}</a:solidFill><a:miter lim="800000"/></a:ln>')
    return "<a:ln><a:noFill/></a:ln>"


def _fill_xml(lyr: dict) -> str:
    return f'<a:solidFill>{_clr(lyr["fill"], lyr.get("opacity", 1.0))}</a:solidFill>' if lyr.get("fill") else "<a:noFill/>"


def shape_xml(u: _Units, lyr: dict) -> str:
    sid = next_id()
    name = attr(lyr["name"])
    if lyr["shape"] == "line":
        x0, y0, x1, y1 = layout.line_points(lyr)
        bx, by, bw, bh = min(x0, x1), min(y0, y1), abs(x1 - x0), abs(y1 - y0)
        xf = _xfrm(u, bx, by, max(bw, 0.01), max(bh, 0.01), flip_h=x1 < x0, flip_v=y1 < y0)
        return (f'<p:cxnSp {XMLNS}><p:nvCxnSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
                f'<p:spPr>{xf}<a:prstGeom prst="line"><a:avLst/></a:prstGeom>{_line_xml(u, lyr)}</p:spPr></p:cxnSp>')
    ox, oy, ow, oh, r = layout.shape_outline(lyr)
    geom = {"rect": "rect", "rounded": "roundRect", "ellipse": "ellipse"}.get(lyr["shape"], "rect")
    av = ""
    if geom == "roundRect":
        av = f'<a:gd name="adj" fmla="val {int(round(min(50000, r / max(1e-6, min(ow, oh)) * 100000)))}"/>'
    return (f'<p:sp {XMLNS}><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr>{_xfrm(u, ox, oy, ow, oh, lyr.get("rotation", 0.0))}<a:prstGeom prst="{geom}"><a:avLst>{av}</a:avLst>'
            f'</a:prstGeom>{_fill_xml(lyr)}{_line_xml(u, lyr)}</p:spPr>'
            f'<p:txBody><a:bodyPr rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>')


def _custgeom(paths_segs, nw: float, nh: float) -> str:
    """One <a:custGeom> in a 100x coordinate space of the graphic's natural size."""
    W, H = max(1, int(round(nw * 100))), max(1, int(round(nh * 100)))

    def p(x, y):
        return f'<a:pt x="{int(round(x * 100))}" y="{int(round(y * 100))}"/>'

    out = []
    for segs in paths_segs:
        for s in segs:
            k = s[0]
            if k == "M":
                out.append(f"<a:moveTo>{p(s[1], s[2])}</a:moveTo>")
            elif k == "L":
                out.append(f"<a:lnTo>{p(s[1], s[2])}</a:lnTo>")
            elif k == "C":
                out.append(f"<a:cubicBezTo>{p(s[1], s[2])}{p(s[3], s[4])}{p(s[5], s[6])}</a:cubicBezTo>")
            elif k == "Q":
                out.append(f"<a:quadBezTo>{p(s[1], s[2])}{p(s[3], s[4])}</a:quadBezTo>")
            elif k == "Z":
                out.append("<a:close/>")
    return (f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
            f'<a:pathLst><a:path w="{W}" h="{H}">{"".join(out)}</a:path></a:pathLst></a:custGeom>')


def vector_xml(u: _Units, lyr: dict) -> str:
    """A group with one freeform per path (vtracer stacks them, later ones on top)."""
    x, y, w, h = lyr["box"]
    nw, nh = lyr.get("natural") or (w, h)
    gid = next_id()
    kids = []
    for i, p in enumerate(lyr["paths"]):
        segs = parse_path(p["d"])
        if not segs:
            continue
        sid = next_id()
        kids.append(f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{attr(lyr["name"])} part {i + 1}"/><p:cNvSpPr/><p:nvPr/>'
                    f'</p:nvSpPr><p:spPr>{_xfrm(u, x, y, w, h)}{_custgeom([segs], nw, nh)}'
                    f'<a:solidFill>{_clr(p["fill"], lyr.get("opacity", 1.0))}</a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')
    grp = (f'<a:xfrm{_rot(lyr)}><a:off x="{u.emu(x)}" y="{u.emu(y)}"/><a:ext cx="{max(1, u.emu(w))}" cy="{max(1, u.emu(h))}"/>'
           f'<a:chOff x="{u.emu(x)}" y="{u.emu(y)}"/><a:chExt cx="{max(1, u.emu(w))}" cy="{max(1, u.emu(h))}"/></a:xfrm>')
    return (f'<p:grpSp {XMLNS}><p:nvGrpSpPr><p:cNvPr id="{gid}" name="{attr(lyr["name"])}"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
            f'<p:grpSpPr>{grp}</p:grpSpPr>{"".join(kids)}</p:grpSp>')


def _rot(lyr: dict) -> str:
    r = lyr.get("rotation", 0.0)
    return f' rot="{int(round(r * 60000)) % 21600000}"' if r else ""


def run_props(u: _Units, family: str, weight: int, italic: bool, size_px: float, color: str, *, underline=False,
              alpha: float = 1.0, tag: str = "a:rPr") -> str:
    face, bold, it = fonts_out.office_name(family, weight, italic)
    sz = max(100, int(round(u.pt(size_px) * 100)))
    attrs = f'lang="en-US" sz="{sz}" b="{1 if bold else 0}" i="{1 if it else 0}"' + (' u="sng"' if underline else "")
    return (f'<{tag} {attrs} dirty="0"><a:solidFill>{_clr(color, alpha)}</a:solidFill>'
            f'<a:latin typeface="{attr(face)}"/><a:ea typeface="{attr(face)}"/><a:cs typeface="{attr(face)}"/></{tag}>')


def text_xml(u: _Units, lyr: dict) -> str:
    st = lyr["style"]
    x, y, w, h = lyr["box"]
    g = layout.text_geometry(lyr)
    size, pitch = g["size"], g["pitch"]
    top = g["baselines"][0] - (pitch - PPTX_FIRST_BASELINE * size)
    # Same centre as the layer's box, so a rotation turns the text about the same point.
    fh = max(2 * (y + h / 2 - top), pitch * len(g["lines"]))
    algn = {"left": "l", "center": "ctr", "right": "r", "justify": "just"}.get(g["align"], "l")
    rtl = ' rtl="1"' if g["rtl"] else ""
    paras = []
    rpr = run_props(u, st["family"], st["weight"], st["italic"], size, st["color"], underline=st.get("underline", False),
                    alpha=lyr.get("opacity", 1.0))
    for t in g["lines"]:
        ppr = (f'<a:pPr algn="{algn}"{rtl}><a:lnSpc><a:spcPts val="{int(round(u.pt(pitch) * 100))}"/></a:lnSpc>'
               f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft><a:buNone/></a:pPr>')
        run = f"<a:r>{rpr}<a:t>{escape(t)}</a:t></a:r>" if t else ""
        end = rpr.replace("<a:rPr ", "<a:endParaRPr ").replace("</a:rPr>", "</a:endParaRPr>")
        paras.append(f"<a:p>{ppr}{run}{end}</a:p>")
    sid = next_id()
    return (f'<p:sp {XMLNS}><p:nvSpPr><p:cNvPr id="{sid}" name="{attr(lyr["name"])}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr>{_xfrm(u, x, top, w, fh, lyr.get("rotation", 0.0))}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
            f'<a:noFill/></p:spPr><p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" rtlCol="0">'
            f'<a:noAutofit/></a:bodyPr><a:lstStyle/>{"".join(paras)}</p:txBody></p:sp>')


def _border(u: _Units, tag: str, color: str, width: float) -> str:
    return (f'<a:{tag} w="{max(1, int(round(u.pt(width) * EMU_PER_PT)))}" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill>{_clr(color)}</a:solidFill><a:prstDash val="solid"/></a:{tag}>')


def table_xml(u: _Units, lyr: dict) -> str:
    """A native table (graphic frame) with spans, fills, per-cell weight and alignment, and rules."""
    st, border = lyr["style"], lyr["border"]
    x, y, w, h = lyr["box"]
    xs, ys = layout.table_grid(lyr)
    pad = lyr.get("padding") or [6, 2, 6, 2]
    rows, cols = len(ys) - 1, len(xs) - 1
    grid: dict[tuple[int, int], dict] = {}
    for c in lyr["cells"]:
        for r in range(c["row"], min(rows, c["row"] + c.get("rowSpan", 1))):
            for k in range(c["col"], min(cols, c["col"] + c.get("colSpan", 1))):
                grid[(r, k)] = c
    trs = []
    for r in range(rows):
        tcs = []
        for k in range(cols):
            c = grid.get((r, k))
            attrs = ""
            if c is not None and (r, k) == (c["row"], c["col"]):
                if c.get("colSpan", 1) > 1:
                    attrs += f' gridSpan="{c["colSpan"]}"'
                if c.get("rowSpan", 1) > 1:
                    attrs += f' rowSpan="{c["rowSpan"]}"'
            elif c is not None:
                attrs += (' hMerge="1"' if k > c["col"] else "") + (' vMerge="1"' if r > c["row"] else "")
            paras = []
            if c is not None and (r, k) == (c["row"], c["col"]) and c["text"]:
                algn = {"left": "l", "center": "ctr", "right": "r"}.get(c.get("align", "left"), "l")
                rpr = run_props(u, st["family"], c.get("weight", 400), c.get("italic", False), st["size"],
                                c.get("color") or st["color"])
                for t in c["text"].split("\n"):
                    paras.append(f'<a:p><a:pPr algn="{algn}"/>' + (f"<a:r>{rpr}<a:t>{escape(t)}</a:t></a:r>" if t else "")
                                 + rpr.replace("<a:rPr ", "<a:endParaRPr ").replace("</a:rPr>", "</a:endParaRPr>") + "</a:p>")
            if not paras:
                paras = ['<a:p><a:endParaRPr lang="en-US" dirty="0"/></a:p>']
            fill = (f'<a:solidFill>{_clr(c["fill"])}</a:solidFill>' if c is not None and c.get("fill") else "<a:noFill/>")
            anchor = {"top": "t", "bottom": "b"}.get((c or {}).get("valign", "middle"), "ctr")
            borders = "".join(_border(u, t, border["color"], border["width"]) for t in ("lnL", "lnR", "lnT", "lnB"))
            tcpr = (f'<a:tcPr marL="{u.emu(pad[0])}" marR="{u.emu(pad[2])}" marT="{u.emu(pad[1])}" marB="{u.emu(pad[3])}" '
                    f'anchor="{anchor}">{borders}{fill}</a:tcPr>')
            tcs.append(f'<a:tc{attrs}><a:txBody><a:bodyPr/><a:lstStyle/>{"".join(paras)}</a:txBody>{tcpr}</a:tc>')
        trs.append(f'<a:tr h="{u.emu(ys[r + 1] - ys[r])}">{"".join(tcs)}</a:tr>')
    grid_cols = "".join(f'<a:gridCol w="{u.emu(b - a)}"/>' for a, b in zip(xs, xs[1:]))
    gid = next_id()
    return (f'<p:graphicFrame {XMLNS}><p:nvGraphicFramePr><p:cNvPr id="{gid}" name="{attr(lyr["name"])}"/>'
            f'<p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
            f'{_xfrm(u, x, y, w, h, tag="p:xfrm")}<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
            f'<a:tbl><a:tblPr firstRow="0" bandRow="0"><a:tableStyleId>{NO_STYLE_TABLE}</a:tableStyleId></a:tblPr>'
            f'<a:tblGrid>{grid_cols}</a:tblGrid>{"".join(trs)}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>')


def _background(slide, u: _Units, scene: dict, project: Path) -> None:
    from lxml import etree
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn

    bg = next((lyr for lyr in scene["layers"] if lyr.get("role") == "background" and lyr.get("visible", True)), None)
    csld = slide._element.find(qn("p:cSld"))
    if bg is None:
        xml = (f'<p:bg {XMLNS}><p:bgPr><a:solidFill>{_clr(scene["page"].get("background", "#ffffff"))}</a:solidFill>'
               f'<a:effectLst/></p:bgPr></p:bg>')
    else:
        _, rid = slide.part.get_or_add_image_part(str(project / bg["asset"]))
        xml = (f'<p:bg {XMLNS}><p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:srcRect/>'
               f'<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
    csld.insert(0, parse_xml(xml))
    etree.cleanup_namespaces(csld)


def write(scene: dict, project: Path, out: Path) -> tuple[Path, list[str]]:
    from pptx import Presentation
    from pptx.oxml import parse_xml
    from pptx.util import Emu

    k, note = layout.page_scale(scene, max_in=56, min_in=1)
    u = _Units(k)
    W, H = scene["page"]["width"], scene["page"]["height"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(u.emu(W)), Emu(u.emu(H))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _background(slide, u, scene, project)
    tree = slide.shapes._spTree
    for lyr in layout.visible_layers(scene):
        t = lyr["type"]
        if lyr.get("role") == "background":
            continue
        x, y, w, h = lyr["box"]
        if t == "image":
            pic = slide.shapes.add_picture(str(project / lyr["asset"]), Emu(u.emu(x)), Emu(u.emu(y)), Emu(u.emu(w)), Emu(u.emu(h)))
            pic.name = lyr["name"]
            if lyr.get("rotation"):
                pic.rotation = float(lyr["rotation"])
            if lyr.get("opacity", 1.0) < 0.999:
                blip = pic._element.blipFill.find("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
                blip.append(parse_xml(f'<a:alphaModFix {XMLNS} amt="{int(lyr["opacity"] * 100000)}"/>'))
        elif t == "shape":
            tree.append(parse_xml(shape_xml(u, lyr)))
        elif t == "vector":
            tree.append(parse_xml(vector_xml(u, lyr)))
        elif t == "text":
            tree.append(parse_xml(text_xml(u, lyr)))
        elif t == "table":
            tree.append(parse_xml(table_xml(u, lyr)))
    prs.core_properties.title = scene["source"].get("name", "")
    prs.save(out)
    return out, [note] if note else []
